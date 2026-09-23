#!/usr/bin/env python3
"""Deterministic CV workspace validation, DOCX rendering, and version diff."""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import importlib
import json
import os
import re
import sys
import tempfile
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import yaml

REPO = Path(__file__).resolve().parent.parent
PROJECTS = REPO / "projects"
SCRIPTS = REPO / ".scripts"
sys.path.insert(0, str(SCRIPTS))
import source_locator as sl

SCHEMA = "cv-function-result-v1"
VERSION_SCHEMA = "cv-version-v1"
TEMPLATE_VERSION = "native-docx-v1"
TEMPLATE_SCHEMA = "cv-template-v1"
VERIFICATION_STATES = {"draft", "needs_verification", "verified", "retired"}
LANGUAGE_KEYS = {"zh-CN": "zh", "zh": "zh", "en": "en", "en-US": "en", "en-GB": "en"}
RAW_DOMAINS = {"academic", "admin", "teaching", "business", "cross-domain", "private"}
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WORD = f"{{{WORD_NS}}}"


class CvError(RuntimeError):
    pass


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _json_hash(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _result(kind: str, status: str = "ok", **values: Any) -> dict[str, Any]:
    return {"schema": SCHEMA, "kind": kind, "status": status, **values}


def resolve_workspace(value: str) -> Path:
    requested = Path(value)
    candidate = requested if requested.is_absolute() else PROJECTS / requested
    resolved = candidate.resolve()
    projects = PROJECTS.resolve()
    try:
        resolved.relative_to(projects)
    except ValueError as exc:
        raise CvError(f"CV workspace must stay under {PROJECTS}") from exc
    if resolved == projects or not resolved.is_dir():
        raise CvError(f"CV workspace does not exist: {value}")
    records = resolved / "cv-records.yaml"
    if not records.is_file() or records.is_symlink():
        raise CvError(f"CV workspace lacks a regular cv-records.yaml: {value}")
    return resolved


def load_workspace(value: str) -> tuple[Path, Path, dict[str, Any]]:
    workspace = resolve_workspace(value)
    source = workspace / "cv-records.yaml"
    try:
        data = yaml.safe_load(source.read_text(encoding="utf-8")) or {}
    except (OSError, yaml.YAMLError) as exc:
        raise CvError(f"Cannot read cv-records.yaml: {exc}") from exc
    if not isinstance(data, dict):
        raise CvError("cv-records.yaml top level must be a mapping")
    return workspace, source, data


def _template_paths(workspace: Path, template_id: str) -> tuple[Path, Path]:
    if not isinstance(template_id, str) or not re.fullmatch(r"[a-z0-9][a-z0-9-]*", template_id):
        raise CvError("edition template must be a lowercase letters/digits/hyphens ID")
    templates = workspace / "templates"
    if not templates.is_dir() or templates.is_symlink():
        raise CvError("workspace templates must be a regular directory")
    spec_path = templates / f"{template_id}.template.yaml"
    docx_path = templates / f"{template_id}.docx"
    for path in (spec_path, docx_path):
        if not path.is_file() or path.is_symlink():
            raise CvError(f"template package file is missing or unsafe: {path.name}")
    return spec_path, docx_path


def _docx_template_shape(path: Path) -> tuple[list[str], set[str]]:
    try:
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
    except (OSError, KeyError, zipfile.BadZipFile, ET.ParseError) as exc:
        raise CvError(f"cannot read template DOCX: {exc}") from exc
    bookmarks = []
    placeholders: set[str] = set()
    for paragraph in root.iter(WORD + "p"):
        text = "".join(item.text or "" for item in paragraph.iter(WORD + "t"))
        placeholders.update(re.findall(r"\{\{[^{}]+\}\}", text))
        for item in paragraph.iter(WORD + "bookmarkStart"):
            name = item.get(WORD + "name", "")
            if name.startswith("cv_module_"):
                bookmarks.append(name.removeprefix("cv_module_"))
    return bookmarks, placeholders


def _validate_template_spec(template_id: str, spec: Any, docx_path: Path) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    if not isinstance(spec, dict):
        return [_issue("error", "template_shape", "template YAML top level must be a mapping")]
    if spec.get("schema") != TEMPLATE_SCHEMA:
        issues.append(_issue("error", "template_schema", f"template schema must be {TEMPLATE_SCHEMA}"))
    if spec.get("id") != template_id:
        issues.append(_issue("error", "template_id", "template YAML id does not match edition template"))
    if spec.get("renderer") != TEMPLATE_VERSION:
        issues.append(_issue("error", "template_renderer", f"template renderer must be {TEMPLATE_VERSION}"))
    modules = spec.get("modules")
    if not isinstance(modules, list) or not modules:
        issues.append(_issue("error", "template_modules", "template modules must be a non-empty list"))
        return issues
    module_ids: list[str] = []
    mapped_categories: set[str] = set()
    required_placeholders: set[str] = set()
    profile_count = 0
    for index, module in enumerate(modules):
        if not isinstance(module, dict):
            issues.append(_issue("error", "template_module_shape", f"template module {index + 1} must be a mapping"))
            continue
        module_id = module.get("id")
        if not isinstance(module_id, str) or not re.fullmatch(r"[a-z][a-z0-9_]*", module_id):
            issues.append(_issue("error", "template_module_id", f"template module {index + 1} has an invalid id"))
            continue
        if module_id in module_ids:
            issues.append(_issue("error", "template_module_duplicate", f"duplicate template module: {module_id}"))
        module_ids.append(module_id)
        kind = module.get("kind", "records")
        if kind == "profile":
            profile_count += 1
            fields = module.get("fields")
            if not isinstance(fields, list) or not fields or not all(isinstance(item, str) for item in fields):
                issues.append(_issue("error", "template_profile_fields", "profile module fields must be a non-empty string list"))
            else:
                unsupported = sorted(set(fields) - {"name", "headline", "research_interests"})
                if unsupported:
                    issues.append(_issue("error", "template_profile_fields", f"unsupported profile fields: {', '.join(unsupported)}"))
                required_placeholders.update(f"{{{{profile.{field}}}}}" for field in fields)
            continue
        if kind != "records":
            issues.append(_issue("error", "template_module_kind", f"unsupported module kind: {kind}"))
            continue
        categories = module.get("categories")
        labels = module.get("label")
        required = module.get("required_categories", [])
        if not isinstance(categories, list) or not categories or not all(isinstance(item, str) and item for item in categories):
            issues.append(_issue("error", "template_categories", f"module {module_id} categories must be a non-empty string list"))
            categories = []
        duplicates = mapped_categories.intersection(categories)
        if duplicates:
            issues.append(_issue("error", "template_category_duplicate", f"categories mapped more than once: {', '.join(sorted(duplicates))}"))
        mapped_categories.update(categories)
        if not isinstance(labels, dict) or not all(isinstance(labels.get(key), str) and labels[key].strip() for key in ("zh", "en")):
            issues.append(_issue("error", "template_label", f"module {module_id} requires zh and en labels"))
        if not isinstance(required, list) or not all(isinstance(item, str) and item in categories for item in required):
            issues.append(_issue("error", "template_required_categories", f"module {module_id} required_categories must be a subset of categories"))
        if module.get("sort", "source") not in {"source", "date_asc", "date_desc"}:
            issues.append(_issue("error", "template_sort", f"module {module_id} has an unsupported sort"))
        required_placeholders.update({f"{{{{module.{module_id}.label}}}}", f"{{{{module.{module_id}.records}}}}"})
    if profile_count != 1:
        issues.append(_issue("error", "template_profile_count", "template must define exactly one profile module"))
    try:
        bookmarks, placeholders = _docx_template_shape(docx_path)
    except CvError as exc:
        issues.append(_issue("error", "template_docx", str(exc)))
        return issues
    if bookmarks != module_ids:
        issues.append(_issue("error", "template_bookmarks", "DOCX module bookmarks must exactly match YAML module order"))
    missing = sorted(required_placeholders - placeholders)
    if missing:
        issues.append(_issue("error", "template_placeholders", f"DOCX is missing placeholders: {', '.join(missing)}"))
    return issues


def load_template(workspace: Path, data: dict[str, Any], edition_id: str) -> tuple[Path, Path, dict[str, Any], list[dict[str, str]]]:
    edition = (data.get("editions") or {}).get(edition_id) or {}
    template_id = edition.get("template")
    try:
        spec_path, docx_path = _template_paths(workspace, template_id)
        spec = yaml.safe_load(spec_path.read_text(encoding="utf-8")) or {}
    except (CvError, OSError, yaml.YAMLError) as exc:
        return Path(), Path(), {}, [_issue("error", "template_load", str(exc), edition=edition_id)]
    return spec_path, docx_path, spec, _validate_template_spec(template_id, spec, docx_path)


def _issue(severity: str, code: str, message: str, record_id: str = "",
           edition: str = "") -> dict[str, str]:
    item = {"severity": severity, "code": code, "message": message}
    if record_id:
        item["record_id"] = record_id
    if edition:
        item["edition"] = edition
    return item


def _raw_path(path_text: str) -> bool:
    parts = Path(path_text).parts
    return "raw" in parts and bool(set(parts) & RAW_DOMAINS)


def _docx_paragraph_count(path: Path) -> int:
    try:
        with zipfile.ZipFile(path) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
    except (OSError, KeyError, zipfile.BadZipFile, ET.ParseError):
        return 0
    body = root.find(f"{{{WORD_NS}}}body")
    return len(body.findall(f"{{{WORD_NS}}}p")) if body is not None else 0


def evidence_status(locator: str, cache: dict[str, tuple[bool, str]] | None = None) -> tuple[bool, str]:
    cache = cache if cache is not None else {}
    if locator in cache:
        return cache[locator]
    path_text, fragment = sl.split_locator(locator)
    if not path_text:
        result = (False, "evidence must include a path")
    elif not _raw_path(path_text):
        result = (False, "evidence path is outside a Raw domain")
    else:
        target = sl.resolve_path(path_text)
        if target is None or not target.is_file():
            result = (False, "evidence path does not exist")
        elif not fragment:
            result = (True, "whole-document")
        elif target.suffix.lower() == ".docx" and re.fullmatch(r"paragraph-([1-9]\d*)", fragment):
            number = int(fragment.split("-", 1)[1])
            count = _docx_paragraph_count(target)
            result = (number <= count, "present" if number <= count else f"paragraph exceeds {count}")
        else:
            read_target = target
            normalized_fragment = re.sub(r"^pages-", "page-", fragment, flags=re.I)
            status = sl.locator_status(normalized_fragment, read_target)
            if status != "present" and target.suffix.lower() in sl.BINARY_SUFFIXES:
                companion = target.with_name(sl.locator_companion_name(target.name))
                if companion.is_file():
                    status = sl.locator_status(normalized_fragment, companion)
            result = (status == "present", status)
    cache[locator] = result
    return result


def validate_structure(data: dict[str, Any]) -> list[dict[str, str]]:
    issues: list[dict[str, str]] = []
    if data.get("schema_version") != 1:
        issues.append(_issue("error", "schema_version", "schema_version must be 1"))
    profile = data.get("profile")
    if not isinstance(profile, dict) or not isinstance(profile.get("name"), dict):
        issues.append(_issue("error", "profile", "profile.name must be a mapping"))
    editions = data.get("editions")
    if not isinstance(editions, dict) or not editions:
        issues.append(_issue("error", "editions", "editions must be a non-empty mapping"))
        editions = {}
    else:
        for edition_id, edition in editions.items():
            if not isinstance(edition, dict):
                issues.append(_issue("error", "edition_shape", "edition must be a mapping", edition=edition_id))
                continue
            if edition.get("language") not in LANGUAGE_KEYS:
                issues.append(_issue("error", "edition_language", "edition language is unsupported", edition=edition_id))
            if not isinstance(edition.get("template"), str) or not edition["template"].strip():
                issues.append(_issue("error", "edition_template", "edition template must be a non-empty ID", edition=edition_id))
    records = data.get("records")
    if not isinstance(records, list):
        issues.append(_issue("error", "records", "records must be a list"))
        return issues
    seen: set[str] = set()
    for index, record in enumerate(records):
        if not isinstance(record, dict):
            issues.append(_issue("error", "record_shape", f"record {index + 1} must be a mapping"))
            continue
        record_id = record.get("id") if isinstance(record.get("id"), str) else ""
        if not record_id:
            issues.append(_issue("error", "record_id", f"record {index + 1} lacks id"))
        elif record_id in seen:
            issues.append(_issue("error", "duplicate_id", "record id is duplicated", record_id))
        seen.add(record_id)
        for field in ("category", "date"):
            if not isinstance(record.get(field), str) or not record[field].strip():
                issues.append(_issue("error", f"record_{field}", f"record requires non-empty {field}", record_id))
        title = record.get("title")
        if not isinstance(title, dict) or not any(isinstance(title.get(key), str) and title[key].strip() for key in ("zh", "en")):
            issues.append(_issue("error", "record_title", "record requires title.zh or title.en", record_id))
        if record.get("verification") not in VERIFICATION_STATES:
            issues.append(_issue("error", "record_verification", "record verification state is invalid", record_id))
        evidence = record.get("evidence")
        if not isinstance(evidence, list) or not all(isinstance(item, str) and item.strip() for item in evidence):
            issues.append(_issue("error", "record_evidence", "record evidence must be a non-empty string list", record_id))
        include_in = record.get("include_in")
        if not isinstance(include_in, list) or not all(isinstance(item, str) and item in editions for item in include_in):
            issues.append(_issue("error", "record_editions", "record include_in contains an undeclared edition", record_id))
    return issues


def selected_records(data: dict[str, Any], edition_id: str) -> list[dict[str, Any]]:
    editions = data.get("editions") or {}
    if edition_id not in editions:
        raise CvError(f"Unknown CV edition: {edition_id}")
    return [record for record in data.get("records", [])
            if isinstance(record, dict) and edition_id in (record.get("include_in") or [])]


def _language(data: dict[str, Any], edition_id: str) -> str:
    language = data["editions"][edition_id].get("language")
    if language not in LANGUAGE_KEYS:
        raise CvError(f"Unsupported edition language: {language}")
    return LANGUAGE_KEYS[language]


def record_text(record: dict[str, Any], language: str) -> tuple[str, bool]:
    title = record.get("title") or {}
    primary = title.get(language)
    if isinstance(primary, str) and primary.strip():
        return primary.strip(), False
    source_language = str(record.get("source_language") or "").lower()
    original_language_category = record.get("category") in {"book", "publication", "conference_paper"}
    if language == "zh" and (source_language.startswith("en") or original_language_category):
        fallback = title.get("en")
        if isinstance(fallback, str) and fallback.strip():
            return fallback.strip(), True
    return "", False


def validate_edition(workspace: Path, data: dict[str, Any], edition_id: str) -> dict[str, Any]:
    all_issues = validate_structure(data)
    selected_ids = {
        str(record.get("id")) for record in data.get("records", [])
        if isinstance(record, dict) and edition_id in (record.get("include_in") or [])
    }
    issues = [item for item in all_issues
              if not item.get("record_id") or item["record_id"] in selected_ids]
    spec_path, docx_path, template, template_issues = load_template(workspace, data, edition_id)
    issues.extend(template_issues)
    template_id = ((data.get("editions") or {}).get(edition_id) or {}).get("template")
    if any(item["severity"] == "error" for item in issues):
        return {"edition": edition_id, "template": template_id, "selected": len(selected_ids),
                "errors": issues, "warnings": []}
    language = _language(data, edition_id)
    records = selected_records(data, edition_id)
    cache: dict[str, tuple[bool, str]] = {}
    formal_issues: list[dict[str, str]] = []
    warnings: list[dict[str, str]] = []
    selected_categories = {str(record.get("category") or "") for record in records}
    record_modules = [module for module in template["modules"] if module.get("kind", "records") == "records"]
    mapped_categories = {category for module in record_modules for category in module["categories"]}
    required_categories = [category for module in record_modules for category in module.get("required_categories", [])]
    for category in required_categories:
        if category not in selected_categories:
            formal_issues.append(_issue("error", "missing_required_category",
                                        f"edition lacks required category: {category}", edition=edition_id))
    for category in sorted(selected_categories - mapped_categories):
        formal_issues.append(_issue("error", "unmapped_category",
                                    f"selected category is not mapped by template: {category}", edition=edition_id))
    for record in records:
        record_id = record["id"]
        if record.get("verification") != "verified":
            formal_issues.append(_issue("error", "not_verified", "selected record is not verified", record_id, edition_id))
        usable = []
        for locator in record.get("evidence") or []:
            ok, detail = evidence_status(locator, cache)
            usable.append(ok)
            if not ok:
                warnings.append(_issue("warning", "evidence_unusable", f"unusable evidence {locator}: {detail}", record_id, edition_id))
        if not any(usable):
            formal_issues.append(_issue("error", "no_usable_evidence", "selected record has no usable Raw locator", record_id, edition_id))
        text, fallback = record_text(record, language)
        if not text:
            formal_issues.append(_issue("error", "missing_language", f"selected record lacks {language} text", record_id, edition_id))
        elif fallback:
            warnings.append(_issue("warning", "source_language_fallback", "Chinese edition uses verified English source text", record_id, edition_id))
    if not records:
        formal_issues.append(_issue("error", "empty_edition", "edition selects no records", edition=edition_id))
    return {
        "edition": edition_id,
        "template": template_id,
        "template_spec": str(spec_path.relative_to(REPO)),
        "template_docx": str(docx_path.relative_to(REPO)),
        "language": language,
        "selected": len(records),
        "errors": formal_issues,
        "warnings": warnings,
    }


def status_report(workspace: Path, source: Path, data: dict[str, Any]) -> dict[str, Any]:
    records = [record for record in data.get("records", []) if isinstance(record, dict)]
    counts = Counter(str(record.get("verification") or "missing") for record in records)
    editions = {}
    for edition_id in (data.get("editions") or {}):
        report = validate_edition(workspace, data, edition_id)
        editions[edition_id] = {
            "language": report.get("language"),
            "template": report.get("template"),
            "selected": report["selected"],
            "errors": len(report["errors"]),
            "warnings": len(report["warnings"]),
            "ready": not report["errors"],
        }
    return _result(
        "status", workspace=str(workspace.relative_to(REPO)),
        source=str(source.relative_to(REPO)), records=len(records),
        verification=dict(sorted(counts.items())), editions=editions,
        workspace_issues=validate_structure(data),
        categories=dict(sorted(Counter(str(record.get("category") or "missing") for record in records).items())),
    )


def _load_docx():
    try:
        return importlib.import_module("docx")
    except ModuleNotFoundError as exc:
        if exc.name != "docx":
            raise CvError(f"python-docx dependency failed: {exc.name}") from exc
    explicit = os.environ.get("CV_DOCX_SITE")
    candidates = [Path(explicit).expanduser().resolve()] if explicit else sorted(
        (Path.home() / ".cache/codex-runtimes").glob("*/dependencies/python/lib/python*/site-packages"),
        reverse=True,
    )
    for candidate in candidates:
        if not (candidate / "docx/__init__.py").is_file():
            continue
        sys.path.insert(0, str(candidate))
        try:
            return importlib.import_module("docx")
        except ImportError as exc:
            raise CvError(f"bundled python-docx could not load: {exc}") from exc
    raise CvError("python-docx==1.2.0 is unavailable; install it or set CV_DOCX_SITE")


def doctor_report(workspace: Path, source: Path) -> dict[str, Any]:
    try:
        module = _load_docx()
        dependency = {"status": "available", "version": getattr(module, "__version__", "unknown"),
                      "module": str(Path(module.__file__).resolve())}
        status = "ready"
    except Exception as exc:
        dependency = {"status": "unavailable", "error": str(exc)}
        status = "unavailable"
    return _result("doctor", status=status, workspace=str(workspace.relative_to(REPO)),
                   source=str(source.relative_to(REPO)), dependency=dependency,
                   renderer=TEMPLATE_VERSION, remote_calls=0)


def _date_key(record: dict[str, Any]) -> tuple[int, int, str]:
    value = str(record.get("date") or "")
    matches = re.findall(r"((?:19|20)\d{2})(?:[-年./](\d{1,2}))?", value)
    if not matches:
        return (0, 0, value)
    year, month = matches[-1]
    return (int(year), int(month or 0), value)


def _module_records(records: list[dict[str, Any]], module: dict[str, Any]) -> list[dict[str, Any]]:
    categories = set(module.get("categories") or [])
    selected = [record for record in records if record.get("category") in categories]
    ordering = module.get("sort", "source")
    if ordering == "date_asc":
        return sorted(selected, key=_date_key)
    if ordering == "date_desc":
        return sorted(selected, key=_date_key, reverse=True)
    return selected


def _set_paragraph_element_text(paragraph: Any, text: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run_properties = None
    for child in paragraph:
        if child.tag == WORD + "r":
            candidate = child.find(WORD + "rPr")
            if candidate is not None:
                run_properties = deepcopy(candidate)
            break
    for child in list(paragraph):
        if child.tag != WORD + "pPr":
            paragraph.remove(child)
    run = OxmlElement("w:r")
    if run_properties is not None:
        run.append(run_properties)
    value = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" "):
        value.set(qn("xml:space"), "preserve")
    value.text = text
    run.append(value)
    paragraph.append(run)


def _remove_paragraph(paragraph: Any) -> None:
    parent = paragraph._p.getparent()
    if parent is not None:
        parent.remove(paragraph._p)


def _token_paragraph(document: Any, token: str) -> Any:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text == token]
    if len(matches) != 1:
        raise CvError(f"template token must occur in exactly one paragraph: {token}")
    return matches[0]


def _profile_value(data: dict[str, Any], field: str, language: str) -> str:
    profile = data.get("profile") or {}
    if field == "name":
        return str((profile.get("name") or {}).get(language) or (profile.get("name") or {}).get("en") or "CV")
    if field == "headline":
        if profile.get("verification") != "verified":
            return ""
        return str((profile.get("headline") or {}).get(language) or "")
    if field == "research_interests":
        interests = profile.get("research_interests") or {}
        values = interests.get(language) if profile.get("research_interests_verification") == "verified" else []
        if not values:
            return ""
        label = "Research Interests" if language == "en" else "研究方向"
        return f"{label}: " + "; ".join(str(item) for item in values)
    raise CvError(f"unsupported profile field: {field}")


def _remove_bookmarks(document: Any) -> None:
    for tag in (WORD + "bookmarkStart", WORD + "bookmarkEnd"):
        for element in list(document._element.iter(tag)):
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)


def build_docx(path: Path, template_docx: Path, template: dict[str, Any],
               data: dict[str, Any], edition_id: str,
               records: list[dict[str, Any]]) -> dict[str, Any]:
    docx = _load_docx()
    document = docx.Document(template_docx)
    language = _language(data, edition_id)
    profile = data.get("profile") or {}
    name = (profile.get("name") or {}).get(language) or (profile.get("name") or {}).get("en") or "CV"
    rendered = 0
    for module in template["modules"]:
        module_id = module["id"]
        if module.get("kind", "records") == "profile":
            for field in module["fields"]:
                paragraph = _token_paragraph(document, f"{{{{profile.{field}}}}}")
                value = _profile_value(data, field, language)
                if value:
                    _set_paragraph_element_text(paragraph._p, value)
                else:
                    _remove_paragraph(paragraph)
            continue
        label_paragraph = _token_paragraph(document, f"{{{{module.{module_id}.label}}}}")
        records_paragraph = _token_paragraph(document, f"{{{{module.{module_id}.records}}}}")
        section_records = _module_records(records, module)
        if not section_records:
            _remove_paragraph(label_paragraph)
            _remove_paragraph(records_paragraph)
            continue
        _set_paragraph_element_text(label_paragraph._p, str(module["label"][language]))
        parent = records_paragraph._p.getparent()
        anchor = records_paragraph._p
        for record in section_records:
            text, _ = record_text(record, language)
            clone = deepcopy(records_paragraph._p)
            _set_paragraph_element_text(clone, text)
            parent.insert(parent.index(anchor), clone)
            rendered += 1
        parent.remove(anchor)

    _remove_bookmarks(document)
    document.core_properties.title = f"{name} - {edition_id}"
    document.core_properties.author = str(name)
    document.core_properties.subject = "Evidence-backed CV generated by WikiGraph"
    document.save(path)
    return {"paragraphs": len(document.paragraphs), "records": rendered}


def _bookmark(paragraph: Any, module_id: str, bookmark_id: int, start: bool) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if start:
        element = OxmlElement("w:bookmarkStart")
        element.set(qn("w:id"), str(bookmark_id))
        element.set(qn("w:name"), f"cv_module_{module_id}")
        position = 1 if len(paragraph) and paragraph[0].tag == WORD + "pPr" else 0
        paragraph.insert(position, element)
    else:
        element = OxmlElement("w:bookmarkEnd")
        element.set(qn("w:id"), str(bookmark_id))
        paragraph.append(element)


def _scrub_template_package(path: Path) -> None:
    relationships_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    content_types_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    blocked_prefixes = ("customXml/", "word/media/", "docProps/thumbnail")
    blocked_names = {"docProps/custom.xml"}
    handle, temporary_name = tempfile.mkstemp(prefix=".cv-template-scrub-", suffix=".docx", dir=path.parent)
    os.close(handle)
    temporary = Path(temporary_name)
    try:
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                name = info.filename
                if name in blocked_names or name.startswith(blocked_prefixes):
                    continue
                if re.fullmatch(r"word/_rels/(?:header|footer)\d+\.xml\.rels", name):
                    continue
                payload = source.read(name)
                if name in {"_rels/.rels", "word/_rels/document.xml.rels"}:
                    root = ET.fromstring(payload)
                    for relation in list(root):
                        target_value = relation.get("Target", "")
                        relation_type = relation.get("Type", "")
                        if ("custom-properties" in relation_type or "customXml" in relation_type
                                or "thumbnail" in relation_type
                                or "image" in relation_type or "hyperlink" in relation_type
                                or target_value.startswith(("media/", "../customXml/", "customXml/",
                                                            "docProps/thumbnail", "thumbnail"))):
                            root.remove(relation)
                    ET.register_namespace("", relationships_ns)
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                elif name == "[Content_Types].xml":
                    root = ET.fromstring(payload)
                    for item in list(root):
                        part_name = item.get("PartName", "").lstrip("/")
                        if part_name in blocked_names or part_name.startswith(blocked_prefixes):
                            root.remove(item)
                    ET.register_namespace("", content_types_ns)
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                elif name == "docProps/app.xml":
                    root = ET.fromstring(payload)
                    for item in root.iter():
                        local_name = item.tag.rsplit("}", 1)[-1]
                        if local_name in {"Manager", "Company", "HyperlinkBase"}:
                            item.text = ""
                        elif local_name in {"TotalTime", "Pages", "Words", "Characters", "Lines",
                                            "Paragraphs", "CharactersWithSpaces"}:
                            item.text = "0"
                    if root.tag.startswith("{"):
                        ET.register_namespace("", root.tag[1:].split("}", 1)[0])
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                elif name == "docProps/core.xml":
                    root = ET.fromstring(payload)
                    for item in list(root):
                        if item.tag.rsplit("}", 1)[-1] == "lastPrinted":
                            root.remove(item)
                    ET.register_namespace("cp", "http://schemas.openxmlformats.org/package/2006/metadata/core-properties")
                    ET.register_namespace("dc", "http://purl.org/dc/elements/1.1/")
                    ET.register_namespace("dcterms", "http://purl.org/dc/terms/")
                    ET.register_namespace("dcmitype", "http://purl.org/dc/dcmitype/")
                    ET.register_namespace("xsi", "http://www.w3.org/2001/XMLSchema-instance")
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                target.writestr(info, payload)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def extract_anonymized_template(source: Path, destination: Path, template: dict[str, Any],
                                paragraph_map: dict[str, int]) -> None:
    """Build a clean placeholder template while reusing paragraph styles and page setup."""
    docx = _load_docx()
    source = source.resolve()
    destination = destination.resolve()
    if not source.is_file() or source.is_symlink():
        raise CvError("template source must be a regular DOCX")
    document = docx.Document(source)
    source_paragraphs = [deepcopy(paragraph._p) for paragraph in document.paragraphs]
    required_tokens = []
    for module in template.get("modules") or []:
        module_id = module["id"]
        if module.get("kind", "records") == "profile":
            required_tokens.extend(f"{{{{profile.{field}}}}}" for field in module["fields"])
        else:
            required_tokens.extend((f"{{{{module.{module_id}.label}}}}",
                                    f"{{{{module.{module_id}.records}}}}"))
    missing = [token for token in required_tokens if token not in paragraph_map]
    if missing:
        raise CvError(f"paragraph map lacks template tokens: {', '.join(missing)}")
    if any(not isinstance(paragraph_map[token], int) or paragraph_map[token] < 0
           or paragraph_map[token] >= len(source_paragraphs) for token in required_tokens):
        raise CvError("paragraph map contains an out-of-range source paragraph")

    body = document._body._element
    section_properties = body.find(WORD + "sectPr")
    for child in list(body):
        if child is not section_properties:
            body.remove(child)
    for section in document.sections:
        for container in (section.header, section.footer):
            element = container._element
            for child in list(element):
                element.remove(child)
            from docx.oxml import OxmlElement
            element.append(OxmlElement("w:p"))

    for bookmark_id, module in enumerate(template["modules"], start=1):
        module_id = module["id"]
        if module.get("kind", "records") == "profile":
            tokens = [f"{{{{profile.{field}}}}}" for field in module["fields"]]
        else:
            tokens = [f"{{{{module.{module_id}.label}}}}", f"{{{{module.{module_id}.records}}}}"]
        inserted = []
        for token in tokens:
            paragraph = deepcopy(source_paragraphs[paragraph_map[token]])
            _set_paragraph_element_text(paragraph, token)
            position = body.index(section_properties) if section_properties is not None else len(body)
            body.insert(position, paragraph)
            inserted.append(paragraph)
        _bookmark(inserted[0], module_id, bookmark_id, True)
        _bookmark(inserted[-1], module_id, bookmark_id, False)

    now = datetime.now().replace(microsecond=0)
    properties = document.core_properties
    properties.title = "Academic CV Template"
    properties.subject = "Anonymized template for deterministic CV rendering"
    properties.author = "WikiGraph"
    properties.last_modified_by = "WikiGraph"
    properties.comments = ""
    properties.keywords = ""
    properties.category = ""
    properties.revision = 1
    properties.created = now
    properties.modified = now
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    _scrub_template_package(destination)


def _safe_basename(data: dict[str, Any], edition_id: str, language: str) -> str:
    configured = (data.get("editions", {}).get(edition_id) or {}).get("filename_prefix")
    if isinstance(configured, str) and configured.strip():
        value = configured.strip()
    else:
        name = (data.get("profile", {}).get("name") or {}).get(language) or edition_id
        value = f"{name}-CV" if language == "en" else f"{name}-个人简历"
    value = re.sub(r"[\\/:*?\"<>|]", "-", value).strip(" .")
    if not value:
        raise CvError("edition filename prefix is empty after sanitization")
    return value


def _version_paths(versions: Path, prefix: str, stamp: str) -> tuple[Path, Path, int]:
    if not re.fullmatch(r"\d{8}", stamp):
        raise CvError("version date must use YYYYMMDD")
    revision = 1
    while True:
        suffix = "" if revision == 1 else f"-r{revision}"
        artifact = versions / f"{prefix}-{stamp}{suffix}.docx"
        manifest = versions / f"{prefix}-{stamp}{suffix}.manifest.json"
        if not artifact.exists() and not manifest.exists():
            return artifact, manifest, revision
        revision += 1


def _write_json_atomic(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(handle, "w", encoding="utf-8") as stream:
            json.dump(value, stream, ensure_ascii=False, indent=2, sort_keys=True)
            stream.write("\n")
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temp_name, path)
    except Exception:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise


def _rebuild_index(versions: Path) -> Path:
    entries = []
    for path in sorted(versions.glob("*.manifest.json")):
        try:
            manifest = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if manifest.get("schema") == VERSION_SCHEMA:
            entries.append({key: manifest.get(key) for key in (
                "version_id", "created_at", "edition", "language", "artifact", "artifact_sha256",
            )})
    target = versions / "index.jsonl"
    handle, temp_name = tempfile.mkstemp(prefix=".index.", suffix=".tmp", dir=versions)
    try:
        with os.fdopen(handle, "w", encoding="utf-8") as stream:
            for entry in entries:
                stream.write(json.dumps(entry, ensure_ascii=False, sort_keys=True) + "\n")
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temp_name, target)
    except Exception:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass
        raise
    return target


def render(workspace: Path, source: Path, data: dict[str, Any], edition_id: str,
           stamp: str | None = None) -> dict[str, Any]:
    report = validate_edition(workspace, data, edition_id)
    if report["errors"]:
        return _result("render", status="blocked", workspace=str(workspace.relative_to(REPO)),
                       edition=edition_id, validation=report,
                       error="formal CV validation failed")
    records = selected_records(data, edition_id)
    template_spec_path, template_docx_path, template, template_issues = load_template(workspace, data, edition_id)
    if template_issues:
        raise CvError("template changed after validation")
    language = report["language"]
    versions = workspace / "versions"
    if versions.exists() and (versions.is_symlink() or not versions.is_dir()):
        raise CvError("versions must be a regular directory inside the workspace")
    versions.mkdir(parents=True, exist_ok=True)
    prefix = _safe_basename(data, edition_id, language)
    version_date = stamp or date.today().strftime("%Y%m%d")
    artifact, manifest_path, revision = _version_paths(versions, prefix, version_date)
    descriptor, temporary_name = tempfile.mkstemp(prefix=".cv-render-", suffix=".docx", dir=versions)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        mechanical = build_docx(temporary, template_docx_path, template, data, edition_id, records)
        with zipfile.ZipFile(temporary) as archive:
            if "word/document.xml" not in archive.namelist():
                raise CvError("generated DOCX lacks word/document.xml")
        artifact_sha = _sha256(temporary)
        record_snapshots = {}
        for record in records:
            text, fallback = record_text(record, language)
            record_snapshots[record["id"]] = {
                "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
                "evidence_sha256": _json_hash(record.get("evidence") or []),
                "fallback": fallback,
            }
        manifest = {
            "schema": VERSION_SCHEMA,
            "version_id": artifact.stem,
            "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            "edition": edition_id,
            "language": language,
            "version_date": version_date,
            "revision": revision,
            "workspace": str(workspace.relative_to(REPO)),
            "source": str(source.relative_to(REPO)),
            "source_sha256": _sha256(source),
            "renderer": TEMPLATE_VERSION,
            "template": template["id"],
            "template_spec": str(template_spec_path.relative_to(REPO)),
            "template_spec_sha256": _sha256(template_spec_path),
            "template_docx": str(template_docx_path.relative_to(REPO)),
            "template_docx_sha256": _sha256(template_docx_path),
            "artifact": str(artifact.relative_to(REPO)),
            "artifact_sha256": artifact_sha,
            "record_ids": [record["id"] for record in records],
            "records": record_snapshots,
            "warnings": report["warnings"],
            "mechanical": mechanical,
        }
        if artifact.exists() or manifest_path.exists():
            raise CvError("version collision detected before publish")
        os.replace(temporary, artifact)
        try:
            _write_json_atomic(manifest_path, manifest)
            index = _rebuild_index(versions)
        except Exception:
            artifact.unlink(missing_ok=True)
            manifest_path.unlink(missing_ok=True)
            raise
    finally:
        temporary.unlink(missing_ok=True)
    return _result("render", workspace=str(workspace.relative_to(REPO)), edition=edition_id,
                   artifact=str(artifact.relative_to(REPO)), manifest=str(manifest_path.relative_to(REPO)),
                   index=str(index.relative_to(REPO)), artifact_sha256=artifact_sha,
                   selected=len(records), warnings=report["warnings"])


def history(workspace: Path) -> dict[str, Any]:
    versions = workspace / "versions"
    entries = []
    if versions.is_dir():
        for path in sorted(versions.glob("*.manifest.json"), reverse=True):
            try:
                item = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                continue
            if item.get("schema") == VERSION_SCHEMA:
                entries.append({key: item.get(key) for key in (
                    "version_id", "created_at", "edition", "language", "artifact", "artifact_sha256",
                )})
    return _result("history", workspace=str(workspace.relative_to(REPO)), versions=entries)


def _manifest_path(workspace: Path, value: str) -> Path:
    versions = (workspace / "versions").resolve()
    requested = Path(value)
    if requested.is_absolute():
        candidate = requested.resolve()
    elif "/" in value or "\\" in value:
        candidate = (REPO / requested).resolve()
    else:
        candidate = (versions / requested).resolve()
    if candidate.suffix == ".docx":
        candidate = candidate.with_suffix(".manifest.json")
    elif not candidate.name.endswith(".manifest.json"):
        candidate = candidate.with_name(candidate.name + ".manifest.json")
    try:
        candidate.relative_to(versions)
    except ValueError as exc:
        raise CvError("version manifest must stay inside the workspace versions directory") from exc
    if not candidate.is_file():
        raise CvError(f"version manifest does not exist: {value}")
    return candidate


def diff_versions(workspace: Path, before: str, after: str) -> dict[str, Any]:
    before_path = _manifest_path(workspace, before)
    after_path = _manifest_path(workspace, after)
    left = json.loads(before_path.read_text(encoding="utf-8"))
    right = json.loads(after_path.read_text(encoding="utf-8"))
    if left.get("schema") != VERSION_SCHEMA or right.get("schema") != VERSION_SCHEMA:
        raise CvError("both inputs must be cv-version-v1 manifests")
    left_records = left.get("records") or {}
    right_records = right.get("records") or {}
    left_ids, right_ids = set(left_records), set(right_records)
    changed = sorted(record_id for record_id in left_ids & right_ids
                     if left_records[record_id] != right_records[record_id])
    return _result(
        "diff", workspace=str(workspace.relative_to(REPO)),
        before=left.get("version_id"), after=right.get("version_id"),
        added=sorted(right_ids - left_ids), removed=sorted(left_ids - right_ids),
        changed=changed, unchanged=len((left_ids & right_ids) - set(changed)),
        edition_changed=left.get("edition") != right.get("edition"),
        source_changed=left.get("source_sha256") != right.get("source_sha256"),
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Evidence-backed CV workspace function")
    sub = parser.add_subparsers(dest="command", required=True)
    def add_workspace(command_parser: argparse.ArgumentParser) -> None:
        command_parser.add_argument("--workspace", default="简历维护", help="projects/ 下的 CV 工作区")

    status = sub.add_parser("status", help="汇总记录与 edition 就绪状态")
    add_workspace(status)
    validate = sub.add_parser("validate", help="校验一个正式 edition")
    validate.add_argument("edition")
    add_workspace(validate)
    render_parser = sub.add_parser("render", help="生成不可覆盖的日期 DOCX 版本")
    render_parser.add_argument("edition")
    render_parser.add_argument("--date", default="", help="版本日期 YYYYMMDD，默认今天")
    render_parser.add_argument("--format", default="docx", choices=("docx",))
    add_workspace(render_parser)
    history_parser = sub.add_parser("history", help="列出带 manifest 的正式版本")
    add_workspace(history_parser)
    diff_parser = sub.add_parser("diff", help="比较两个版本 manifest")
    diff_parser.add_argument("before")
    diff_parser.add_argument("after")
    add_workspace(diff_parser)
    doctor_parser = sub.add_parser("doctor", help="检查本地 DOCX 渲染依赖")
    add_workspace(doctor_parser)
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        workspace, source, data = load_workspace(args.workspace)
        if args.command == "status":
            result = status_report(workspace, source, data)
        elif args.command == "validate":
            report = validate_edition(workspace, data, args.edition)
            result = _result("validate", status="blocked" if report["errors"] else "ok",
                             workspace=str(workspace.relative_to(REPO)), **report)
        elif args.command == "render":
            result = render(workspace, source, data, args.edition, args.date or None)
        elif args.command == "history":
            result = history(workspace)
        elif args.command == "diff":
            result = diff_versions(workspace, args.before, args.after)
        elif args.command == "doctor":
            result = doctor_report(workspace, source)
        else:
            raise CvError(f"unknown command: {args.command}")
    except Exception as exc:
        result = _result(args.command, status="error", error=f"{type(exc).__name__}: {exc}"[:2000])
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0 if result.get("status") == "ok" or args.command in {"status", "history", "diff", "doctor"} and result.get("status") != "error" else 1


if __name__ == "__main__":
    raise SystemExit(main())
