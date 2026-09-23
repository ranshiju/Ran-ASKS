#!/usr/bin/env python3
"""Shared native slide/deck compiler. No remote QA or user approvals."""
from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import subprocess
import xml.etree.ElementTree as ET
import zipfile

VERSION = 'presentation-single-slide-v1'
CANVAS = [13.333, 7.5]


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                     separators=(',', ':'), allow_nan=False).encode()).hexdigest()


def file_hash(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def fingerprint():
    """Hash local renderer and font inventory, never serialize the environment."""
    import presentation_runtime as runtime
    import visual_qa as qa
    info = runtime.doctor()
    if info['status'] != 'ready':
        raise ValueError('Local rendering dependencies unavailable')
    soffice = qa._find_soffice()
    env = qa._soffice_env(soffice)
    fc = shutil.which('fc-list')
    if not fc:
        raise ValueError('fc-list required to fingerprint local font files')
    proc = subprocess.run([fc, '--format=%{file}\n'], env=env, capture_output=True,
                          text=True, timeout=30, check=True)
    font_paths = sorted(set(proc.stdout.splitlines()))
    if not font_paths:
        raise ValueError('Font inventory is empty; cannot bind a render snapshot')
    fonts = [(str(Path(p).resolve()), file_hash(p)) for p in font_paths]
    config = env.get('FONTCONFIG_FILE')
    # Version + source/binary/font hashes; no credentials or full environment.
    return {'version': VERSION, 'compiler_sha256': file_hash(__file__),
            'qa_sha256': file_hash(qa.__file__), 'runtime_sha256': file_hash(runtime.__file__),
            'dependencies': info['dependencies'], 'soffice_sha256': file_hash(soffice),
            'soffice_version': subprocess.run([str(soffice), '--version'], env=env,
                capture_output=True, text=True, timeout=30, check=True).stdout.strip(),
            'font_inventory_sha256': digest(fonts),
            'fontconfig_sha256': file_hash(config) if config else None}


def validate_ir(ir):
    if not isinstance(ir, dict) or set(ir) != {'schema', 'slide_id', 'canvas', 'objects'}:
        raise ValueError('Invalid IR fields')
    if ir['schema'] != 'presentation-ir-v1' or ir['canvas'] != CANVAS:
        raise ValueError('Unsupported IR version/canvas')
    objects = ir['objects']
    if not isinstance(objects, list) or not 1 <= len(objects) <= 30:
        raise ValueError('Invalid object count')
    ids = set()
    for obj in objects:
        if not isinstance(obj, dict) or not isinstance(obj.get('id'), str) or not re.fullmatch(r'[a-zA-Z0-9][a-zA-Z0-9_-]{0,63}', obj['id']):
            raise ValueError('Object ID required')
        if obj['id'] in ids:
            raise ValueError('Duplicate object ID')
        ids.add(obj['id'])
        kind = obj.get('type')
        extras = {'text': {'text', 'size'}, 'shape': {'text', 'size'},
                  'image': {'asset'}, 'connector': {'start', 'end'}}
        if kind not in extras or set(obj) != {'id', 'type', 'bbox'} | extras[kind]:
            raise ValueError('Unsupported IR object/fields')
        box = obj['bbox']
        if (not isinstance(box, list) or len(box) != 4
                or any(type(x) not in (int, float) or not math.isfinite(x) for x in box)):
            raise ValueError('Invalid geometry')
        x, y, w, h = box
        if min(x, y, w, h) < 0 or w <= 0 or (h <= 0 and kind != 'connector'):
            raise ValueError('Invalid geometry dimensions')
        if x + w > CANVAS[0] + 1e-6 or y + h > CANVAS[1] + 1e-6:
            raise ValueError('Object outside canvas')
        if kind in ('text', 'shape'):
            if not isinstance(obj['text'], str) or not obj['text'].strip():
                raise ValueError('Empty native text')
            if type(obj['size']) not in (int, float) or not 18 <= obj['size'] <= 44:
                raise ValueError('Unsupported font size')
        if kind == 'image' and (not isinstance(obj['asset'], str) or not re.fullmatch('[a-f0-9]{64}', obj['asset'])):
            raise ValueError('Image asset hash required')
    shapes = {x['id'] for x in objects if x['type'] == 'shape'}
    for obj in objects:
        if obj['type'] == 'connector' and (obj['start'] not in shapes or obj['end'] not in shapes
                                         or obj['start'] == obj['end']):
            raise ValueError('Dangling connector')
    return ir


def compile_slide(content, design):
    objects = [{'id': 'title', 'type': 'text', 'bbox': [.7, .55, 11.9, 1.1],
                'text': content['title'], 'size': 32}]
    body = content['body']
    layout = design['layout_id']
    if layout == 'title':
        objects.append({'id': 'body', 'type': 'text', 'bbox': [.9, 2.1, 11.5, 4.5],
                        'text': '\n'.join(body), 'size': 26})
    elif layout == 'image_text':
        source = next(x for x in content['sources'] if x['id'] == design['image_source'])
        objects += [{'id': 'image', 'type': 'image', 'bbox': [.7, 2, 6, 4.6],
                     'asset': source['sha256']},
                    {'id': 'body', 'type': 'text', 'bbox': [7.1, 2, 5.5, 4.6],
                     'text': '\n'.join(body), 'size': 24}]
    elif layout == 'flow':
        if not 2 <= len(body) <= 4:
            raise ValueError('Flow requires 2–4 body items')
        width, gap = (11.9 - .5 * (len(body) - 1)) / len(body), .5
        for i, text in enumerate(body):
            x = .7 + i * (width + gap)
            objects.append({'id': f'node{i}', 'type': 'shape', 'bbox': [x, 3, width, 1.6],
                            'text': text, 'size': 24})
            if i:
                objects.append({'id': f'edge{i}', 'type': 'connector',
                                'bbox': [x - gap, 3.8, gap, 0],
                                'start': f'node{i-1}', 'end': f'node{i}'})
    else:
        raise ValueError('Unsupported layout')
    return validate_ir({'schema': 'presentation-ir-v1', 'slide_id': content['slide_id'],
                        'canvas': CANVAS, 'objects': objects})


def _native_verify(path, ir, page=1):
    ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
          'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read(f'ppt/slides/slide{page}.xml'))
    expected = {kind: sum(x['type'] == kind for x in ir['objects'])
                for kind in ('text', 'shape', 'connector', 'image')}
    if (len(root.findall('.//p:sp', ns)) != expected['text'] + expected['shape']
            or len(root.findall('.//p:pic', ns)) != expected['image']
            or len(root.findall('.//p:cxnSp', ns)) != expected['connector']):
        raise ValueError('Native object counts differ from compiled IR')
    shape_names = {}
    expected_text = {x['id']: x['text'] for x in ir['objects'] if x['type'] in ('text', 'shape')}
    actual_text = {}
    for shape in root.findall('.//p:sp', ns):
        identity = shape.find('p:nvSpPr/p:cNvPr', ns)
        shape_names[identity.get('id')] = identity.get('name')
        actual_text[identity.get('name')] = '\n'.join(
            ''.join(node.text or '' for node in paragraph.findall('.//a:t', ns))
            for paragraph in shape.findall('p:txBody/a:p', ns))
    if actual_text != expected_text:
        raise ValueError('Serialized native text/identity differs from IR (including page order)')
    ids = set(shape_names)
    for connector in root.findall('.//p:cxnSp', ns):
        endpoints = [connector.find(f'p:nvCxnSpPr/p:cNvCxnSpPr/a:{tag}', ns)
                     for tag in ('stCxn', 'endCxn')]
        if any(x is None or x.get('id') not in ids for x in endpoints):
            raise ValueError('Native connector lost its bound endpoint')
        name = connector.find('p:nvCxnSpPr/p:cNvPr', ns).get('name')
        expected_connector = next((x for x in ir['objects'] if x['id'] == name and x['type'] == 'connector'), None)
        if expected_connector is None or [shape_names[x.get('id')] for x in endpoints] != [expected_connector['start'], expected_connector['end']]:
            raise ValueError('Native connector binds to the wrong shapes')
    return expected


def append_slide(deck, ir, content, theme, assets):
    """Append only validated native IR; never clone uncontrolled slide binaries."""
    from PIL import Image
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    validate_ir(ir)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(theme['background'])
    native = {}
    for obj in ir['objects']:
        x, y, w, h = obj['bbox']
        kind = obj['type']
        if kind == 'connector':
            continue
        if kind == 'image':
            path = assets / obj['asset']
            if file_hash(path) != obj['asset']:
                raise ValueError('Asset snapshot corrupted')
            with Image.open(path) as image:
                iw, ih = image.size
            scale = min(w / iw, h / ih)
            shape = slide.shapes.add_picture(str(path), Inches(x + (w-iw*scale)/2),
                Inches(y + (h-ih*scale)/2), width=Inches(iw*scale), height=Inches(ih*scale))
        else:
            if kind == 'shape':
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *[Inches(v) for v in obj['bbox']])
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(theme['accent'])
            else:
                shape = slide.shapes.add_textbox(*[Inches(v) for v in obj['bbox']])
            shape.text_frame.word_wrap = True
            for index, text in enumerate(obj['text'].split('\n')):
                paragraph = shape.text_frame.paragraphs[0] if index == 0 else shape.text_frame.add_paragraph()
                run = paragraph.add_run()
                run.text = text
                run.font.name = theme['font']
                run.font.size = Pt(obj['size'])
                run.font.color.rgb = RGBColor.from_string(theme['foreground'])
        shape.name = obj['id']
        native[obj['id']] = shape
    for obj in ir['objects']:
        if obj['type'] != 'connector':
            continue
        x, y, w, h = obj['bbox']
        shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
            Inches(x), Inches(y), Inches(x+w), Inches(y+h))
        shape.begin_connect(native[obj['start']], 3)
        shape.end_connect(native[obj['end']], 1)
        shape.name = obj['id']
    slide.notes_slide.notes_text_frame.text = content['notes']
    return slide


def render(content, design, theme, assets, output):
    """Render actual local artifacts; callers cannot supply a pass receipt."""
    import presentation_runtime as runtime
    import visual_qa as qa
    import fitz
    before = fingerprint()
    runtime._load_pptx()
    from pptx import Presentation
    from pptx.util import Inches

    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    ir = compile_slide(content, design)
    deck = Presentation()
    deck.slide_width, deck.slide_height = [Inches(v) for v in CANVAS]
    append_slide(deck, ir, content, theme, assets)
    pptx = output / 'slide.pptx'
    deck.save(pptx)
    counts = _native_verify(pptx, ir)
    kind, pdf, count = qa._prepare_source(pptx, output)
    if count != 1:
        raise ValueError('Single-slide render returned wrong page count')
    png = output / 'preview.png'
    qa.render_page(kind, pdf, 1, png)
    with fitz.open(pdf) as doc:
        actual = ''.join(doc[0].get_text().split())
        for obj in ir['objects']:
            if obj['type'] in ('text', 'shape') and ''.join(obj['text'].split()) not in actual:
                raise ValueError(f'PDF text missing or clipped: {obj["id"]}')
        fonts = sorted({font[3] for font in doc[0].get_fonts(full=True)})
    if fingerprint() != before:
        raise ValueError('Renderer/font dependencies changed during build')
    return ir, {'schema': 'presentation-build-v1', 'renderer': before, 'native_objects': counts,
                'artifacts': {name: {'sha256': file_hash(path)} for name, path in
                              (('pptx', pptx), ('pdf', pdf), ('preview', png))},
                'structural_check': 'passed', 'text_roundtrip': 'passed',
                'requested_font': theme['font'], 'observed_pdf_fonts': fonts,
                'warnings': ['Font equivalence and text overflow require actual visual review'],
                'visual_review': 'not_checked', 'user_approval': 'not_requested',
                'remote_calls': 0}, {'pptx': pptx, 'pdf': pdf, 'preview': png}


def assemble(slides, theme, assets, output):
    """Generate a deck and previews. Mechanical checks are NOT form/evidence review."""
    import presentation_runtime as runtime
    import visual_qa as qa
    import fitz
    before = fingerprint()
    runtime._load_pptx()
    from pptx import Presentation
    from pptx.util import Inches
    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    deck = Presentation()
    deck.slide_width, deck.slide_height = [Inches(v) for v in CANVAS]
    for item in slides:
        append_slide(deck, item['ir'], item['content'], theme, assets)
    pptx = output / 'deck.pptx'
    deck.save(pptx)
    counts = [_native_verify(pptx, item['ir'], i + 1) for i, item in enumerate(slides)]
    # Opening the saved native package and rendering ensure it is usable, not that it looks right.
    if len(Presentation(pptx).slides) != len(slides):
        raise ValueError('Native deck page count mismatch')
    kind, pdf, count = qa._prepare_source(pptx, output)
    if count != len(slides):
        raise ValueError('Rendered deck page count mismatch')
    paths = {'pptx': pptx, 'pdf': pdf}
    fonts = []
    with fitz.open(pdf) as doc:
        for index, item in enumerate(slides):
            png = output / f'page-{index + 1}.png'
            qa.render_page(kind, pdf, index + 1, png)
            paths[f"preview:{item['ir']['slide_id']}"] = png
            fonts.append(sorted({f[3] for f in doc[index].get_fonts(full=True)}))
    if fingerprint() != before:
        raise ValueError('Renderer/font dependencies changed during deck build')
    return {'renderer': before, 'page_count': count, 'native_objects': counts,
            'structural_check': 'passed', 'requested_font': theme['font'],
            'observed_pdf_fonts': fonts, 'visual_review': 'not_checked',
            'user_approval': 'not_requested', 'remote_calls': 0,
            'warnings': ['Final appearance and PowerPoint editing require user inspection']}, paths
