#!/usr/bin/env python3
"""Isolated collaboration tests; synthetic confirmations are never real approval."""
from __future__ import annotations

import argparse
from copy import deepcopy
import json
from pathlib import Path
import tempfile
import subprocess
import sys
from types import SimpleNamespace
import unittest
from unittest.mock import patch

import science_exhibit as core
import science_exhibit_render as renderer
import science_exhibit_workflow as work
from test_science_exhibit import content, concept_content, composition_content

REAL_RENDER = False
TEST_INSTRUCTION = 'TEST-ONLY 合成确认，不对应真实用户指令或科学批准'


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name).resolve()
        self.source = self.root / 'academic/raw/references/fixture/paper.md'
        self.source.parent.mkdir(parents=True)
        self.source.write_text('测试模型仅用于工程检查，不对应真实科学研究。\n第二行：这是用于验证文本引用的合成材料。\n', encoding='utf-8')
        (self.root / 'projects/fixture').mkdir(parents=True)
        self.repo_patch = patch.object(core, 'REPO', self.root)
        self.repo_patch.start()
        self.addCleanup(self.repo_patch.stop)
        self.store = self.root / 'projects/fixture/exhibits/collaboration'
        core.init_store(self.store)

    def state(self):
        return core.load_state(self.store)

    def action(self, action, **changes):
        options = {'action': action, 'store': self.store, 'expected_revision': self.state()['revision']}
        options.update(changes)
        return work.act(core, SimpleNamespace(**options))

    def candidate(self, kind, value=None, *, commit=True):
        task = core.prepare(self.store, self.state()['revision'], kind=kind)
        candidate = deepcopy(value if value is not None else task['agent_task']['protocol']['output_template'])
        output = self.root / task['agent_task']['outputs'][0]['path']
        output.write_text(json.dumps(candidate, ensure_ascii=False), encoding='utf-8')
        request = self.root / task['request_path']
        result = core.check_or_commit(request, commit=commit)
        return result, request, candidate

    def settle_brief(self):
        requirements = work.initial_state()['requirements']
        for key in work.REQUIRED:
            requirements[key] = {'value': 'TEST-ONLY ' + key, 'status': 'confirmed', 'basis': TEST_INSTRUCTION}
        brief = {'schema': 'science-exhibit-brief-v1', 'source_path': str(self.source.relative_to(self.root)),
                 'source_label': '仅供工程测试的合成材料，不对应真实论文', 'requirements': requirements}
        self.candidate('brief', brief)

    def plan(self, kind='flow'):
        return {
            'schema': 'science-exhibit-plan-v1', 'title': 'TEST-ONLY 科学展板工程检查',
            'goal': '检查协作协议，不代表科学结论', 'viewer_question': '怎样保持版本和证据一致？',
            'story': '问题、机械方法、检查结果与边界',
            'visual': {'kind': kind, 'description': '原生可编辑工程流程示意，不是观测图'},
            'layout': '中文 A1 竖版，四个文字区域和解释性示意', 'canvas_mm': [594, 841],
            'claims': content(self.source, self.root)['claims'],
            'boundaries': ['仅用于工程检查，不对应真实科学研究'], 'open_questions': [],
        }

    def authorize(self):
        self.settle_brief()
        self.candidate('plan', self.plan())
        binding = work.binding(core, self.state())
        self.action('confirm-plan', binding_sha256=binding, instruction=TEST_INSTRUCTION)
        self.action('start-making', binding_sha256=binding, instruction=TEST_INSTRUCTION)

    def commit_content(self):
        self.candidate('content', content(self.source, self.root))

    def export(self):
        def fake_render(content_value, brief, directory, *, canvas_mm=None):
            for name in work.ARTIFACTS:
                (directory / name).write_bytes(b'TEST-ONLY synthetic rendered artifact')
            return {'mechanical_status': 'passed', 'remote_calls': 0, 'canvas_mm': canvas_mm,
                    'renderer_sha256': renderer.compiler_hash(), 'font_sha256': 'a' * 64}

        with patch.object(renderer, 'render', side_effect=fake_render):
            result = core.build(self.store, self.state()['revision'])
        self.assertIn('record_preview', result['commands'])
        self.action('record-preview', output_dir=result['output_dir'])
        return result

    def test_composition_plan_and_content_bind_to_authorized_story(self):
        self.settle_brief()
        self.candidate('plan', self.plan('composition'))
        binding = work.binding(core, self.state())
        self.action('confirm-plan', binding_sha256=binding, instruction=TEST_INSTRUCTION)
        self.action('start-making', binding_sha256=binding, instruction=TEST_INSTRUCTION)
        candidate = composition_content(self.root, self.source, self.store)
        self.candidate('content', candidate)
        self.assertEqual(self.state()['content']['schema'], 'science-exhibit-content-v2')
        self.assertEqual(core.status(self.store)['phase'], 'making')
        with self.assertRaises(ValueError):
            self.candidate('content', content(self.source, self.root))

    def test_review_brief_adapts_to_nonphysics_audiences_without_story_answers(self):
        import review
        import science_exhibit_review as integration
        output = self.store / 'outputs/test-only'
        output.mkdir(parents=True)
        args = SimpleNamespace(brief=None, reviewer='reader-comprehension', host_capabilities='TEST-ONLY',
            sensitivity='public', allow_host_processing=True, accept_unenforced_tools=True, require_enforced_tools=False)
        scenarios = [
            ('TEST-ONLY 高中生', '生态观察展板，三分钟浏览，不使用类比'),
            ('TEST-ONLY 社区成人读者', '语言学习研究展板，十分钟浏览，按问题和发现组织'),
        ]
        for audience, setting in scenarios:
            with self.subTest(audience=audience):
                state = self.state()
                state['workflow']['requirements']['audience']['value'] = audience
                state['workflow']['requirements']['setting']['value'] = setting
                state['workflow']['requirements']['goal']['value'] = 'TEST-ONLY 作者预期答案，不交给盲读者'
                with patch.object(review, 'REPO', self.root), \
                        patch.object(work, 'validate_preview', return_value={'output_dir': str(output.relative_to(self.root))}), \
                        patch.object(review, 'prepare', return_value={}) as prepared:
                    integration.prepare_review(core, self.store, state, args)
                brief = core.read_json(prepared.call_args.args[1])
                self.assertEqual(brief['audience'], audience)
                self.assertEqual(brief['use_context'], setting)
                self.assertNotIn('作者预期答案', str(brief))
                for fixed_story in ('MPE', '深海', '潜艇', '五分钟'):
                    self.assertNotIn(fixed_story, str(brief))

    def test_review_regions_forward_to_shared_core_with_bound_pdf(self):
        import review
        import science_exhibit_review as integration
        output = self.store / 'outputs/test-only-views'
        output.mkdir(parents=True)
        args = SimpleNamespace(brief=None, reviewer='reader-comprehension', host_capabilities='TEST-ONLY',
            sensitivity='public', allow_host_processing=True, accept_unenforced_tools=True,
            require_enforced_tools=False, image_regions='temp/TEST-ONLY-regions.json')
        with patch.object(review, 'REPO', self.root), \
                patch.object(work, 'validate_preview', return_value={'output_dir': str(output.relative_to(self.root))}), \
                patch.object(review, 'prepare', return_value={}) as prepared:
            integration.prepare_review(core, self.store, self.state(), args)
        self.assertEqual(prepared.call_args.kwargs['image_regions'], args.image_regions)
        self.assertEqual(prepared.call_args.kwargs['detail_source'], output / 'board.pdf')
        self.assertEqual(prepared.call_args.kwargs['artifact_manifest'], output / 'manifest.json')

    def test_image_review_decisions_comparison_and_stale_application(self):
        import shutil
        import review
        import science_exhibit_review as integration
        from test_review import SOURCE, feedback
        from PIL import Image
        self.test_composition_plan_and_content_bind_to_authorized_story()

        def fake_render(content_value, brief, directory, *, canvas_mm=None):
            for name in ('board.pptx', 'board.pdf'):
                (directory / name).write_bytes(b'TEST-ONLY fake output, no actual rendering')
            Image.new('RGB', (320, 480), '#267F86').save(directory / 'preview.png')
            core.atomic_json(directory / 'layout.json', {'objects': [
                {'id': item['id'], 'type': 'text', 'text': item['text']} for item in content_value['texts']]})
            return {'mechanical_status': 'passed', 'remote_calls': 0, 'canvas_mm': canvas_mm,
                    'renderer_sha256': renderer.compiler_hash(), 'font_sha256': 'a' * 64}

        with patch.object(review, 'REPO', self.root), patch.object(renderer, 'render', side_effect=fake_render), \
                patch.object(renderer, 'font_info', return_value={'sha256': 'a' * 64}):
            shutil.copytree(SOURCE / 'agents/reviewers', self.root / 'agents/reviewers')
            first = core.build(self.store, self.state()['revision'])
            host = self.root / 'temp/host.json'
            review.write_json(host, {'schema': 'review-host-capabilities-v1', 'host_id': 'TEST-ONLY',
                'fresh_context': 'observed', 'tools': 'broad', 'tools_enforced': False, 'processing_location': 'unknown',
                'evidence': 'Synthetic fixture, not an actual host observation', 'limitations': ['No real agent used.']})
            args = SimpleNamespace(brief=None, reviewer='reader-comprehension', host_capabilities=host,
                sensitivity='public', allow_host_processing=True, accept_unenforced_tools=True, require_enforced_tools=False)
            view = integration.prepare_review(core, self.store, self.state(), args)
            run = self.root / view['run']
            self.assertNotIn('claims', str(review.packet(run)))
            review.start(run, 'TEST-ONLY-synthetic-host-agent-call', image_attached=True)
            candidate = feedback()
            candidate['material_observation'] = {'image_received': True, 'readability': 'readable'}
            review.write_json(review.candidate_path(run.name), candidate)
            review.collect(run)
            decisions = self.root / 'temp/decisions.json'
            core.atomic_json(decisions, [{'feedback_id': 'note-1', 'action': 'partly-adopt', 'reason': TEST_INSTRUCTION}])
            record = SimpleNamespace(run=run, decisions=decisions, instruction=TEST_INSTRUCTION)
            result = integration.record_review(core, self.store, self.state(), record)
            self.assertFalse(result['already_recorded'])
            self.assertTrue(integration.record_review(core, self.store, self.state(), record)['already_recorded'])
            self.assertEqual(core.status(self.store)['user_preview_acceptance'], 'not_recorded')
            revised = deepcopy(self.state()['content'])
            revised['texts'][1]['text'] = '怎样让主线更容易理解？'
            self.candidate('content', revised)
            second = core.build(self.store, self.state()['revision'])
            with self.assertRaises(ValueError):
                integration.record_review(core, self.store, self.state(), record)
            comparison = integration.compare(core, self.store, self.state(), first['output_dir'], second['output_dir'])
            self.assertTrue((self.root / comparison['comparison']).is_file())
            self.assertIn('hook', comparison['changed_object_ids'])
            self.assertEqual(integration.compare(core, self.store, self.state(), first['output_dir'], second['output_dir'])['output_dir'],
                             comparison['output_dir'])
            (self.root / first['output_dir'] / 'preview.png').write_bytes(b'corrupt')
            with self.assertRaises(ValueError):
                integration.compare(core, self.store, self.state(), first['output_dir'], second['output_dir'])

    def accept(self):
        with patch.object(renderer, 'font_info', return_value={'sha256': 'a' * 64}):
            return self.action('accept-preview', preview_sha256=self.state()['workflow']['preview']['manifest_sha256'], instruction=TEST_INSTRUCTION)

    def test_empty_start_and_resume_are_read_only(self):
        pointer_hash = core.file_hash(self.store / 'current.json')
        result = core.status(self.store)
        self.assertEqual(result['phase'], 'briefing')
        self.assertFalse(result['source_selected'])
        self.assertEqual(result['missing_requirements'], list(work.REQUIRED))
        self.assertFalse(result['plan_confirmed'])
        self.assertEqual(result['next_actions'][0]['action'], 'prepare')
        self.assertEqual(core.file_hash(self.store / 'current.json'), pointer_hash)
        with self.assertRaises(ValueError):
            core.prepare(self.store, 0, kind='content')
        with self.assertRaises(ValueError):
            core.build(self.store, 0)

    def test_proposals_do_not_become_confirmation(self):
        requirements = work.initial_state('候选题目', '建议受众')['requirements']
        requirements['setting'] = {'value': '开放日', 'status': 'assumption', 'basis': 'Agent 假设'}
        brief = {'schema': 'science-exhibit-brief-v1', 'source_path': str(self.source.relative_to(self.root)),
                 'source_label': 'TEST-ONLY', 'requirements': requirements}
        self.candidate('brief', brief)
        self.candidate('plan', self.plan())
        with self.assertRaises(ValueError):
            self.action('confirm-plan', binding_sha256=work.binding(core, self.state()), instruction=TEST_INSTRUCTION)
        self.assertEqual(core.status(self.store)['phase'], 'briefing')

    def test_brief_validation_stale_requests_and_idempotency(self):
        task = core.prepare(self.store, 0, kind='brief')
        self.assertEqual(task['agent_task']['context']['authoring_method'],
                         'operations/SCIENCE_EXHIBIT.md#公众叙事制作方法')
        template = task['agent_task']['protocol']['output_template']
        template['requirements']['goal'] = {'value': '未确定', 'status': 'unknown', 'basis': ''}
        output = self.root / task['agent_task']['outputs'][0]['path']
        output.write_text(json.dumps(template), encoding='utf-8')
        with self.assertRaises(ValueError):
            core.check_or_commit(self.root / task['request_path'])
        self.settle_brief()
        with self.assertRaises(ValueError):
            core.check_or_commit(self.root / task['request_path'], commit=True)
        result, request, candidate = self.candidate('brief')
        self.assertTrue(core.check_or_commit(request, commit=True)['already_committed'])
        candidate['source_label'] = '改后的候选'
        (request.parent / 'candidate.json').write_text(json.dumps(candidate), encoding='utf-8')
        with self.assertRaises(ValueError):
            core.check_or_commit(request, commit=True)

    def test_plan_quote_open_questions_and_binding(self):
        self.settle_brief()
        plan = self.plan()
        plan['claims'][0]['quote'] = '不存在的原文'
        with self.assertRaises(ValueError):
            self.candidate('plan', plan)
        plan = self.plan()
        plan['open_questions'] = ['TEST-ONLY 尚未确定主视觉']
        self.candidate('plan', plan)
        with self.assertRaises(ValueError):
            self.action('confirm-plan', binding_sha256=work.binding(core, self.state()), instruction=TEST_INSTRUCTION)
        self.candidate('plan', self.plan())
        with self.assertRaises(ValueError):
            self.action('confirm-plan', binding_sha256='0' * 64, instruction=TEST_INSTRUCTION)
        with self.assertRaises(ValueError):
            self.action('start-making', binding_sha256=work.binding(core, self.state()), instruction=TEST_INSTRUCTION)

    def test_unsupported_canvas_is_not_silently_downgraded(self):
        self.settle_brief()
        plan = self.plan('concept')
        plan['canvas_mm'] = [841, 594]
        self.candidate('plan', plan)
        binding = work.binding(core, self.state())
        self.action('confirm-plan', binding_sha256=binding, instruction=TEST_INSTRUCTION)
        self.assertTrue(core.status(self.store)['production_blockers'])
        with self.assertRaises(ValueError):
            self.action('start-making', binding_sha256=binding, instruction=TEST_INSTRUCTION)
        self.assertEqual(self.state()['workflow']['plan']['visual']['kind'], 'concept')

    def test_a0_layout_scales_geometry_and_type_without_reflow(self):
        value = concept_content(self.source, self.root)
        with patch.object(renderer, 'wrap_text', side_effect=lambda text, *args: text):
            nominal = renderer.compile_layout(value, self.state()['brief'], {})
            larger = renderer.compile_layout(value, self.state()['brief'], {}, [841, 1189])
            with self.assertRaises(ValueError):
                renderer.compile_layout(value, self.state()['brief'], {}, [841, 594])
        self.assertEqual(larger['canvas_mm'], [841, 1189])
        factor = min(841 / 594, 1189 / 841)
        for original, enlarged in zip(nominal['objects'], larger['objects']):
            self.assertEqual(original['id'], enlarged['id'])
            if 'bbox_mm' in original:
                for dimension in (2, 3):
                    self.assertAlmostEqual(enlarged['bbox_mm'][dimension], original['bbox_mm'][dimension] * factor)
            if original['type'] == 'text':
                self.assertEqual(original['text'], enlarged['text'])
                self.assertAlmostEqual(enlarged['font_pt'], original['font_pt'] * factor)

    def test_real_a0_flow_and_concept_render(self):
        if not REAL_RENDER:
            self.skipTest('Use --render for actual local rendering')
        for kind in ('flow', 'concept'):
            with self.subTest(kind=kind):
                self.settle_brief()
                plan = self.plan(kind)
                plan['canvas_mm'] = [841, 1189]
                self.candidate('plan', plan)
                binding = work.binding(core, self.state())
                self.action('confirm-plan', binding_sha256=binding, instruction=TEST_INSTRUCTION)
                self.action('start-making', binding_sha256=binding, instruction=TEST_INSTRUCTION)
                value = concept_content(self.source, self.root) if kind == 'concept' else content(self.source, self.root)
                if kind == 'concept':
                    value['title'] = '工程测试样板\n两行问题标题'
                self.candidate('content', value)
                result = core.build(self.store, self.state()['revision'])
                self.assertEqual(result['checks']['canvas_mm'], [841, 1189])
                self.assertEqual(result['checks']['mechanical_status'], 'passed')
                self.assertEqual(result['checks']['remote_calls'], 0)
                from pptx import Presentation
                deck = Presentation(self.root / result['output_dir'] / 'board.pptx')
                from pptx.oxml.ns import qn
                self.assertTrue(all(shape.element.find(qn('p:style')) is None for shape in deck.slides[0].shapes))
                self.assertAlmostEqual(deck.slide_width.mm, 841, places=3)
                self.assertAlmostEqual(deck.slide_height.mm, 1189, places=3)
                self.assertEqual(core.status(self.store)['phase'], 'reviewing')
                self.accept_actual_preview(result)

    def test_concept_reuses_artwork_validation_and_plan_kind(self):
        self.settle_brief()
        self.candidate('plan', self.plan('concept'))
        binding = work.binding(core, self.state())
        self.action('confirm-plan', binding_sha256=binding, instruction=TEST_INSTRUCTION)
        self.assertFalse(core.status(self.store)['production_blockers'])
        self.action('start-making', binding_sha256=binding, instruction=TEST_INSTRUCTION)
        task = core.prepare(self.store, self.state()['revision'], kind='content')
        self.assertEqual(task['agent_task']['protocol']['output_template']['diagram']['kind'], 'concept')
        with self.assertRaisesRegex(ValueError, 'Content visual differs'):
            self.commit_content()
        valid = concept_content(self.source, self.root)
        for update in (
            lambda value: value['diagram'].update(claim_ids=['unknown']),
            lambda value: value['diagram']['artwork']['objects'][0].update(type='image', src='https://example.invalid'),
            lambda value: value['diagram']['artwork']['objects'][0].update(x=-1),
            lambda value: value['diagram']['artwork']['objects'][0].update(fill='url(secret)'),
            lambda value: value['diagram']['artwork'].update(canvas=[1, 720]),
            lambda value: value['diagram']['artwork'].update(objects=[]),
        ):
            invalid = deepcopy(valid)
            update(invalid)
            with self.subTest(update=update), self.assertRaises(ValueError):
                self.candidate('content', invalid)
        self.candidate('content', valid)
        self.assertEqual(self.state()['content']['diagram'], valid['diagram'])

    def test_concept_cannot_replace_authorized_flow(self):
        self.authorize()
        with self.assertRaisesRegex(ValueError, 'Content visual differs'):
            self.candidate('content', concept_content(self.source, self.root))

    def test_artwork_validator_change_invalidates_preview_acceptance(self):
        self.authorize()
        self.commit_content()
        self.export()
        original = renderer.file_hash
        import presentation_artwork
        with patch.object(renderer, 'file_hash', side_effect=lambda path: 'b' * 64 if str(path) == presentation_artwork.__file__ else original(path)):
            with self.assertRaisesRegex(ValueError, 'artwork validator changed'):
                self.accept()

    def test_real_concept_render_and_capacity(self):
        if not REAL_RENDER:
            self.skipTest('Use --render for actual local rendering')
        self.settle_brief()
        self.candidate('plan', self.plan('concept'))
        binding = work.binding(core, self.state())
        self.action('confirm-plan', binding_sha256=binding, instruction=TEST_INSTRUCTION)
        self.action('start-making', binding_sha256=binding, instruction=TEST_INSTRUCTION)
        valid = concept_content(self.source, self.root)
        font = renderer.font_info(self.state()['brief']['font'])
        layout = renderer.compile_layout(valid, self.state()['brief'], font)
        self.assertEqual(layout['visual_kind'], 'concept')
        artwork_objects = [value for value in layout['objects'] if value['id'].startswith('artwork-')]
        self.assertTrue(all(value['claim_ids'] == ['claim-1'] for value in artwork_objects))
        self.assertEqual(next(value for value in artwork_objects if value['id'] == 'artwork-background')['bbox_mm'][2], 538)
        for update in (
            lambda value: value['diagram']['artwork']['objects'][4].update(size=12),
            lambda value: value['diagram']['artwork']['objects'][4].update(w=100),
            lambda value: value['panels'][0].update(body='不能静默缩小字体。' * 40),
        ):
            invalid = deepcopy(valid)
            update(invalid)
            with self.subTest(update=update), self.assertRaises(ValueError):
                renderer.compile_layout(invalid, self.state()['brief'], font)
        self.candidate('content', valid)
        result = core.build(self.store, self.state()['revision'])
        self.assertEqual(result['checks']['mechanical_status'], 'passed')
        self.assertEqual(result['checks']['native_object_count'], len(layout['objects']))
        from pptx import Presentation
        deck = Presentation(self.root / result['output_dir'] / 'board.pptx')
        shapes = {shape.name: shape for shape in deck.slides[0].shapes}
        self.assertEqual(str(shapes['artwork-1'].auto_shape_type), 'OVAL (9)')
        self.assertGreater(shapes['artwork-3'].begin_x, shapes['artwork-3'].end_x)
        self.assertEqual(shapes['artwork-4'].text, '仅供工程测试的图形主视觉检查')
        self.assertEqual(core.status(self.store)['phase'], 'reviewing')
        self.accept_actual_preview(result)

    def accept_actual_preview(self, result):
        self.action('record-preview', output_dir=result['output_dir'])
        preview_sha = self.state()['workflow']['preview']['manifest_sha256']
        self.action('accept-preview', preview_sha256=preview_sha, instruction=TEST_INSTRUCTION)

    def test_lifecycle_separates_authorization_acceptance_and_release(self):
        self.authorize()
        self.assertEqual(core.status(self.store)['phase'], 'making')
        self.assertEqual(core.status(self.store)['user_preview_acceptance'], 'not_recorded')
        self.commit_content()
        self.export()
        self.assertEqual(core.status(self.store)['phase'], 'reviewing')
        with self.assertRaises(ValueError), patch.object(renderer, 'font_info', return_value={'sha256': 'a' * 64}):
            self.action('deliver', preview_sha256=self.state()['workflow']['preview']['manifest_sha256'])
        self.accept()
        with patch.object(renderer, 'font_info', return_value={'sha256': 'a' * 64}):
            result = self.action('deliver', preview_sha256=self.state()['workflow']['preview']['manifest_sha256'])
        self.assertEqual(result['phase'], 'delivered')
        self.assertEqual(result['semantic_review'], 'pending')
        self.assertEqual(result['public_release'], 'not_approved')

    def test_brief_and_plan_changes_invalidate_dependencies(self):
        self.authorize()
        self.commit_content()
        self.export()
        self.accept()
        task = core.prepare(self.store, self.state()['revision'], kind='brief')
        candidate = task['agent_task']['protocol']['output_template']
        candidate['requirements']['audience']['value'] = 'TEST-ONLY 新受众'
        self.candidate('brief', candidate)
        state = self.state()
        self.assertIsNone(state['workflow']['authorization'])
        self.assertIsNone(state['workflow']['preview'])
        self.assertIsNone(state['workflow']['acceptance'])
        self.assertIsNotNone(state['content'])
        binding = work.binding(core, state)
        self.action('confirm-plan', binding_sha256=binding, instruction=TEST_INSTRUCTION)
        self.action('start-making', binding_sha256=binding, instruction=TEST_INSTRUCTION)
        with self.assertRaises(ValueError):
            core.prepare(self.store, self.state()['revision'], panel='finding')
        with self.assertRaises(ValueError):
            core.build(self.store, self.state()['revision'])
        self.commit_content()
        changed = self.plan()
        changed['story'] = '修改后的 TEST-ONLY 科学故事'
        self.candidate('plan', changed)
        self.assertIsNone(self.state()['workflow']['authorization'])

    def test_undefined_visible_acronym_cannot_commit(self):
        self.authorize()
        before = core.file_hash(self.store / 'current.json')
        invalid = content(self.source, self.root)
        invalid['hook'] = 'MPE 尚未定义'
        with self.assertRaisesRegex(ValueError, 'hook: MPE'):
            self.candidate('content', invalid)
        self.assertEqual(core.file_hash(self.store / 'current.json'), before)
        valid = content(self.source, self.root)
        valid['hook'] = '矩阵乘积纠缠（MPE）仅用于工程测试'
        self.candidate('content', valid)

    def test_content_evidence_and_audience_cannot_bypass_plan(self):
        self.authorize()
        changed = content(self.source, self.root)
        changed['claims'][0]['qualifiers'] = []
        with self.assertRaises(ValueError):
            self.candidate('content', changed)
        with self.assertRaises(ValueError):
            core.prepare(self.store, self.state()['revision'], audience='另一个受众')

    def test_feedback_is_persistent_but_not_acceptance(self):
        self.authorize()
        self.commit_content()
        self.export()
        self.accept()
        result = self.action('feedback', scope='visual', text='TEST-ONLY 字太密', instruction=TEST_INSTRUCTION)
        self.assertEqual(result['feedback'][-1]['text'], 'TEST-ONLY 字太密')
        self.assertEqual(result['user_preview_acceptance'], 'not_recorded')
        self.assertTrue(result['production_authorized'])
        result = self.action('feedback', scope='plan', text='TEST-ONLY 改故事', instruction=TEST_INSTRUCTION)
        self.assertEqual(result['phase'], 'planning')
        self.assertFalse(result['production_authorized'])

    def test_content_edit_and_new_export_invalidate_actual_acceptance(self):
        self.authorize()
        self.commit_content()
        first = self.export()
        self.accept()
        changed = content(self.source, self.root)
        changed['panels'][2]['body'] = '本地机械检查不能代替公众理解测试。'
        self.candidate('content', changed)
        self.assertIsNone(self.state()['workflow']['preview'])
        self.assertTrue((self.root / first['output_dir'] / 'board.pptx').is_file())
        self.export()
        self.assertIsNone(self.state()['workflow']['acceptance'])

    def test_artifact_tampering_and_stale_preview_are_rejected(self):
        self.authorize()
        self.commit_content()
        result = self.export()
        with self.assertRaises(ValueError), patch.object(renderer, 'font_info', return_value={'sha256': 'a' * 64}):
            self.action('accept-preview', preview_sha256='0' * 64, instruction=TEST_INSTRUCTION)
        (self.root / result['output_dir'] / 'board.pdf').write_bytes(b'TEST-ONLY tampered')
        with self.assertRaises(ValueError):
            self.accept()
        with self.assertRaises(ValueError):
            core.status(self.store)

    def test_source_drift_and_source_change(self):
        self.authorize()
        self.commit_content()
        second = self.source.parent / 'second.md'
        second.write_text('TEST-ONLY 不同来源\n', encoding='utf-8')
        task = core.prepare(self.store, self.state()['revision'], kind='brief')
        brief = task['agent_task']['protocol']['output_template']
        brief['source_path'] = str(second.relative_to(self.root))
        self.candidate('brief', brief)
        self.assertIsNone(self.state()['content'])
        self.assertIsNone(self.state()['workflow']['plan'])
        second.write_text('TEST-ONLY 漂移\n', encoding='utf-8')
        with self.assertRaises(ValueError):
            core.status(self.store)

    def test_legacy_adoption_does_not_infer_approvals(self):
        legacy = self.root / 'projects/fixture/exhibits/legacy'
        core.init_store(legacy, str(self.source.relative_to(self.root)), 'TEST-ONLY', legacy=True)
        self.assertEqual(core.status(legacy)['phase'], 'legacy_draft')
        self.store = legacy
        result = self.action('adopt')
        self.assertEqual(result['phase'], 'briefing')
        self.assertFalse(result['plan_confirmed'])
        self.assertFalse(result['production_authorized'])
        self.assertEqual(result['user_preview_acceptance'], 'not_recorded')

    def test_failed_candidate_publish_keeps_pointer(self):
        before = core.file_hash(self.store / 'current.json')
        original = core.atomic_json

        def fail_pointer(path, value):
            if Path(path) == self.store / 'current.json':
                raise OSError('TEST-ONLY pointer failure')
            original(path, value)

        with patch.object(core, 'atomic_json', side_effect=fail_pointer), self.assertRaises(OSError):
            self.settle_brief()
        self.assertEqual(core.file_hash(self.store / 'current.json'), before)
        self.assertEqual(self.state()['revision'], 0)

    def test_candidate_json_and_symlink_rejections(self):
        task = core.prepare(self.store, 0, kind='brief')
        request = self.root / task['request_path']
        output = self.root / task['agent_task']['outputs'][0]['path']
        output.write_text('{"schema":"x", "schema":"y"}', encoding='utf-8')
        with self.assertRaises(ValueError):
            core.check_or_commit(request)
        output.unlink()
        output.symlink_to(self.source)
        with self.assertRaises(ValueError):
            core.check_or_commit(request)
        request.write_text('[]', encoding='utf-8')
        with self.assertRaises(ValueError):
            core.check_or_commit(request)

    def test_directory_fsync_failure_after_pointer_replace_rolls_back(self):
        before = core.file_hash(self.store / 'current.json')
        original = core.os.fsync
        failed = False

        def fail_directory_once(descriptor):
            nonlocal failed
            if not failed and core.os.fstat(descriptor).st_ino == self.store.stat().st_ino:
                failed = True
                raise OSError('TEST-ONLY fsync failure after pointer replacement')
            original(descriptor)

        with patch.object(core.os, 'fsync', side_effect=fail_directory_once), self.assertRaises(OSError):
            self.settle_brief()
        self.assertEqual(core.file_hash(self.store / 'current.json'), before)
        self.assertEqual(self.state()['revision'], 0)

    def test_real_collaborative_render(self):
        if not REAL_RENDER:
            self.skipTest('Use --render for actual local rendering')
        self.authorize()
        self.commit_content()
        before = self.state()['revision']
        result = core.build(self.store, before)
        self.assertEqual(self.state()['revision'], before + 1)
        self.assertEqual(core.status(self.store)['phase'], 'reviewing')
        self.assertEqual(result['output_revision'], before)
        self.assertEqual(result['checks']['mechanical_status'], 'passed')
        self.action('record-preview', output_dir=result['output_dir'])
        preview_sha = self.state()['workflow']['preview']['manifest_sha256']
        self.action('accept-preview', preview_sha256=preview_sha, instruction=TEST_INSTRUCTION)
        result = self.action('deliver', preview_sha256=preview_sha)
        self.assertEqual(result['phase'], 'delivered')
        self.assertEqual(result['public_release'], 'not_approved')

    def test_actual_cli_resume_uses_the_same_authority(self):
        script = 'import sys; from pathlib import Path; sys.path.insert(0, sys.argv[1]); import science_exhibit as core; core.REPO=Path(sys.argv[2]); raise SystemExit(core.main(sys.argv[3:]))'
        result = subprocess.run([sys.executable, '-c', script, str(Path(core.__file__).parent), str(self.root),
                                 'resume', '--store', str(self.store.relative_to(self.root))], capture_output=True, text=True, check=True)
        data = json.loads(result.stdout)
        self.assertEqual(data['phase'], 'briefing')
        self.assertFalse(data['production_authorized'])
        self.assertEqual(data['revision'], 0)

    def test_renderer_and_font_changes_block_preview_acceptance(self):
        self.authorize()
        self.commit_content()
        self.export()
        with patch.object(renderer, 'font_info', return_value={'sha256': 'b' * 64}), self.assertRaises(ValueError):
            self.action('accept-preview', preview_sha256=self.state()['workflow']['preview']['manifest_sha256'], instruction=TEST_INSTRUCTION)
        with patch.object(renderer, 'file_hash', return_value='c' * 64), self.assertRaises(ValueError):
            self.accept()

    def test_handwritten_build_receipt_cannot_register_a_preview(self):
        self.authorize()
        self.commit_content()
        state = self.state()
        directory = self.store / 'outputs' / ('r000005-' + 'd' * 32)
        directory.mkdir(parents=True)
        for name in work.ARTIFACTS:
            (directory / name).write_bytes(b'TEST-ONLY handwritten artifact, never built')
        manifest = {'schema': 'science-exhibit-delivery-v1', 'workflow_binding_sha256': work.binding(core, state),
                    'source_sha256': state['source']['sha256'], 'content_sha256': core.digest(state['content']),
                    'files': {name: core.file_hash(directory / name) for name in work.ARTIFACTS},
                    'checks': {'mechanical_status': 'passed', 'remote_calls': 0}}
        core.atomic_json(directory / 'manifest.json', manifest)
        with self.assertRaises(ValueError):
            self.action('record-preview', output_dir=str(directory.relative_to(self.root)))

    def test_build_registration_failure_restores_pointer_and_removes_output(self):
        self.authorize()
        self.commit_content()
        before = core.file_hash(self.store / 'current.json')
        original = core.os.fsync
        failed = False

        def fail_directory_once(descriptor):
            nonlocal failed
            if not failed and core.os.fstat(descriptor).st_ino == self.store.stat().st_ino:
                failed = True
                raise OSError('TEST-ONLY build registration pointer fsync failure')
            original(descriptor)

        with patch.object(core.os, 'fsync', side_effect=fail_directory_once), self.assertRaises(OSError):
            self.export()
        self.assertEqual(core.file_hash(self.store / 'current.json'), before)
        self.assertEqual(list((self.store / 'outputs').iterdir()), [])
        self.assertEqual(core.status(self.store)['phase'], 'making')

    def test_collaboration_route_profiles_load_actual_sections(self):
        route = Path(core.__file__).with_name('route.py')
        for profile in ('discuss', 'resume'):
            result = subprocess.run([sys.executable, str(route), '--capability', 'science-exhibit',
                                     '--capability-profile', profile], capture_output=True, text=True, check=True)
            self.assertNotIn('未命中 section', result.stderr)
            self.assertIn('需求讨论 → 方案讨论', result.stdout)
            if profile == 'resume':
                self.assertIn('## 接续与修改', result.stdout)
        result = subprocess.run([sys.executable, str(route), '--task', 'write'], capture_output=True, text=True, check=True)
        self.assertIn('--capability science-exhibit --capability-profile discuss', result.stdout)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--render', action='store_true')
    args = parser.parse_args()
    REAL_RENDER = args.render
    unittest.main(argv=['test_science_exhibit_workflow.py'], verbosity=2)
