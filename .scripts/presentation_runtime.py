#!/usr/bin/env python3
"""Shared, local-only phase-0 probe; NOT a presentation authoring renderer."""
from __future__ import annotations

import argparse
import hashlib
import importlib
import json
import os
from pathlib import Path
import sys
import tempfile
import zipfile
from xml.etree import ElementTree as ET

REPO = Path(__file__).resolve().parent.parent
SCHEMA = "presentation-runtime-v1"
TEXT_EN = "Editable presentation runtime"
TEXT_ZH = "中文排版测试：知识与证据"
FONTS = {"latin": "Arial", "cjk": "PingFang SC"}


class RuntimeProbeError(RuntimeError):
    """Dependency, path, or mechanical verification failure."""


def _load_pptx():
    """Reuse the existing installed/bundled convention without importing reconstruction."""
    try:
        return importlib.import_module("pptx")
    except ModuleNotFoundError as exc:
        if exc.name != "pptx":
            raise RuntimeProbeError(f"Installed python-pptx dependency missing: {exc.name}") from exc
    explicit = os.environ.get("PRESENTATION_PPTX_SITE")
    if explicit:
        candidates = [Path(explicit).expanduser().resolve()]
    else:
        candidates = sorted((Path.home() / ".cache/codex-runtimes").glob(
            "*/dependencies/python/lib/python*/site-packages"), reverse=True)
    for candidate in candidates:
        if not (candidate / "pptx/__init__.py").is_file():
            continue
        sys.path.insert(0, str(candidate))
        try:
            return importlib.import_module("pptx")
        except ImportError as exc:
            raise RuntimeProbeError(f"Bundled python-pptx could not load: {exc}") from exc
        finally:
            sys.path.pop(0)
    raise RuntimeProbeError(
        "python-pptx unavailable; provision it in the current environment or set "
        "PRESENTATION_PPTX_SITE to a trusted bundled site-packages directory")


def doctor() -> dict:
    """Inspect dependencies only; never load .env or invoke a remote adapter."""
    dependencies = {}
    # Load host dependencies first, so bundled paths cannot shadow their runtime.
    for name in ("PIL", "fitz", "pptx"):
        try:
            module = _load_pptx() if name == "pptx" else importlib.import_module(name)
            dependencies[name] = {
                "status": "available", "version": getattr(module, "__version__", "unknown"),
                "module": str(Path(module.__file__).resolve()),
            }
        except Exception as exc:
            dependencies[name] = {"status": "unavailable", "error": str(exc)[:1500]}
    try:
        qa = importlib.import_module("visual_qa")
        dependencies["soffice"] = {"status": "available", "path": str(qa._find_soffice())}
    except Exception as exc:
        dependencies["soffice"] = {"status": "unavailable", "error": str(exc)[:1500]}
    return {
        "schema": SCHEMA, "kind": "doctor",
        "status": "ready" if all(x["status"] == "available" for x in dependencies.values()) else "unavailable",
        "python": sys.executable, "dependencies": dependencies,
        "scope": "dependency_presence_only", "remote_calls": 0,
        "authoring_available": False,
    }


def _new_run() -> Path:
    root = REPO / "temp" / "presentation-runtime"
    # A fixed output root still needs to reject links in its ancestry.
    if root.resolve() != root:
        raise RuntimeProbeError("presentation runtime cache root/parents must not be symlinks")
    root.mkdir(parents=True, exist_ok=True)
    return Path(tempfile.mkdtemp(prefix="probe-", dir=root))


def _make_sample(output: Path) -> None:
    """Fixed synthetic geometry only; no user input or generic layout interface."""
    pptx = _load_pptx()
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt

    deck = pptx.Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("F5F7FA")

    def label(text, x, y, w, h, size, font, color="16253D"):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.text_frame.word_wrap = False
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.name, run.font.size = font, Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        return shape

    label(TEXT_EN, .7, .65, 12, .7, 32, FONTS["latin"])
    label(TEXT_ZH, .7, 1.55, 12, .7, 28, FONTS["cjk"])
    boxes = []
    for x, text in ((1.1, "Native objects"), (8.2, "Local preview")):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(3.1), Inches(4), Inches(1.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string("E0EBFA")
        shape.line.color.rgb = RGBColor.from_string("3265A8")
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        shape.text = text
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name, run.font.size = FONTS["latin"], Pt(24)
                run.font.color.rgb = RGBColor.from_string("16253D")
        boxes.append(shape)
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
        Inches(5.1), Inches(3.7), Inches(8.2), Inches(3.7))
    connector.begin_connect(boxes[0], 3)
    connector.end_connect(boxes[1], 1)
    connector.line.color.rgb = RGBColor.from_string("3265A8")
    connector.line.width = Pt(2)
    label("Synthetic fixture only - no scientific claims or user approval", .7, 6.25, 12, .5, 18, FONTS["latin"], "53647A")
    deck.core_properties.title = "Presentation phase-0 synthetic runtime probe"
    deck.core_properties.author = "WikiGraph"
    deck.save(output)


def _native_check(path: Path) -> dict:
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with zipfile.ZipFile(path) as archive:
        xml = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
    shapes = xml.findall(".//p:sp", ns)
    connectors = xml.findall(".//p:cxnSp", ns)
    text = "\n".join(node.text or "" for node in xml.findall(".//a:t", ns))
    ids = {node.get("id") for node in xml.findall(".//p:cNvPr", ns)}
    endpoints = [node.get("id") for tag in ("stCxn", "endCxn")
                 for node in xml.findall(f".//a:{tag}", ns)]
    textbox_count = sum(shape.find("p:nvSpPr/p:cNvPr", ns) is not None
                        and shape.find("p:nvSpPr/p:cNvSpPr", ns).get("txBox") == "1"
                        for shape in shapes)
    images = len(xml.findall(".//p:pic", ns))
    if (len(shapes) != 5 or textbox_count != 3 or len(connectors) != 1
            or len(endpoints) != 2 or not set(endpoints) <= ids or images
            or TEXT_EN not in text or TEXT_ZH not in text):
        raise RuntimeProbeError("Synthetic native object/text/connector verification failed")
    return {"textboxes": textbox_count, "shapes": len(shapes) - textbox_count,
            "connectors": len(connectors), "images": images, "bound_endpoints": len(endpoints)}


def _render_sample(source: Path, directory: Path) -> dict:
    import fitz
    import visual_qa
    kind, pdf, count = visual_qa._prepare_source(source, directory)
    if count != 1:
        raise RuntimeProbeError(f"Expected one rendered page, got {count}")
    image = directory / "preview.png"
    visual_qa.render_page(kind, pdf, 1, image, dpi=120)
    with fitz.open(pdf) as document:
        page = document[0]
        text = page.get_text()
        compact = "".join(text.split())
        if any("".join(expected.split()) not in compact for expected in (TEXT_EN, TEXT_ZH)):
            raise RuntimeProbeError("Rendered PDF missing expected English/Chinese text")
        fonts = sorted({span["font"] for block in page.get_text("dict")["blocks"]
                        for line in block.get("lines", []) for span in line["spans"]})
    return {"pdf": str(pdf), "preview": str(image), "page_count": count,
            "text_roundtrip": "pass", "requested_fonts": FONTS,
            "observed_pdf_fonts": fonts, "font_equivalence": "not_asserted"}


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def smoke() -> dict:
    readiness = doctor()
    result = {"schema": SCHEMA, "kind": "smoke", "status": "failed", "stage": "doctor",
              "runtime": readiness, "remote_calls": 0, "authoring_available": False,
              "visual_review": "not_checked", "user_approval": "not_requested"}
    if readiness["status"] != "ready":
        result["error"] = "Local dependencies unavailable; no render was attempted"
        return result
    directory = None
    try:
        result["stage"] = "output_guard"
        directory = _new_run()
        result["run_dir"] = str(directory)
        result["probe_sha256"] = _hash(Path(__file__))
        qa_module = importlib.import_module("visual_qa")
        result["local_renderer_sha256"] = _hash(Path(qa_module.__file__))
        result["stage"] = "native_generation"
        source = directory / "synthetic.pptx"
        _make_sample(source)
        result["stage"] = "native_verification"
        result["native_objects"] = _native_check(source)
        result["stage"] = "local_render"
        result["render"] = _render_sample(source, directory)
        result["artifacts"] = {
            name: {"path": str(path), "sha256": _hash(path)}
            for name, path in (("pptx", source), ("pdf", Path(result["render"]["pdf"])),
                               ("preview", Path(result["render"]["preview"])))
        }
        result.update(status="passed", stage="complete", scope="native_generation_and_local_render_only")
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"[:2000]
    if directory is not None:
        receipt = directory / "receipt.json"
        result["receipt"] = str(receipt)
        pending = directory / "receipt.tmp"
        pending.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        pending.replace(receipt)
    return result


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("doctor", "smoke"))
    args = parser.parse_args(argv)
    try:
        result = doctor() if args.command == "doctor" else smoke()
    except Exception as exc:
        result = {"schema": SCHEMA, "kind": args.command, "status": "failed", "error": str(exc)[:2000]}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] in {"ready", "passed"} else 1


if __name__ == "__main__":
    raise SystemExit(main())
