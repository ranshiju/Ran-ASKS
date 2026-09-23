#!/usr/bin/env python3
"""Isolated mechanical delivery tests; --render uses real native/PDF/PNG artifacts.
All simulated approvals below are TEST-ONLY, never real user acceptance.
"""
from __future__ import annotations
import argparse
from copy import deepcopy
import multiprocessing as mp
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

import presentation_delivery as pd
import presentation_state as ps
import presentation_render as render
from test_presentation_state import StateTests, content, design, request, approval, DECLARATION

REPO = Path(__file__).resolve().parent.parent


def fake_assemble(items, theme, assets, output):
    output.mkdir(parents=True)
    paths = {}
    for name, filename in [('pptx', 'deck.pptx'), ('pdf', 'deck.pdf')] + [
            ('preview:' + x['ir']['slide_id'], x['ir']['slide_id'] + '.png') for x in items]:
        path = output / filename
        path.write_text('TEST-ONLY ' + name + render.digest(items))
        paths[name] = path
    return {'renderer': {'test_only': True}, 'page_count': len(items), 'structural_check': 'passed',
            'visual_review': 'not_checked', 'user_approval': 'not_requested', 'remote_calls': 0}, paths


def child_commit(repo, root, result, queue=None, gate=None, crash=None):
    if crash:
        pd._checkpoint = lambda name: os._exit(73) if name == crash else None
    if gate:
        gate.wait()
    try:
        pd.Delivery(ps.Store(root, repo)).apply(result)
        if queue:
            queue.put('ok')
    except Exception as exc:
        if queue:
            queue.put(str(exc))
        else:
            raise


class DeliveryTests(unittest.TestCase):
    # Share fixture helpers, not a second copy of production state-machine logic.
    apply, ready, locked, source = StateTests.apply, StateTests.ready, StateTests.locked, StateTests.source
    tearDown = StateTests.tearDown

    def setUp(self):
        StateTests.setUp(self)
        self.patches.append(patch.object(render, 'assemble', side_effect=fake_assemble))
        self.assembler = self.patches[-1].start()
        self.delivery = pd.Delivery(self.store)

    def rev(self):
        return self.store.read()['session']['revision_id']

    def exported(self):
        for sid in ['s1', 's2', 's3']:
            self.locked(sid)
        return self.delivery.export(self.rev())

    def candidate(self, kind, *, output=None, slides=None, question=None):
        task = self.delivery.prepare(kind, self.rev(), slides or ['s1'], DECLARATION,
                    output_id=output['output_id'] if output else None, question=question)
        result_path = self.repo / task['outputs'][0]['path']
        result = ps.load_json(result_path)
        result.update(actor='TEST-ONLY host', summary='TEST-ONLY mechanical fixture, not real verification')
        return task, result

    def form(self, output):
        task, result = self.candidate('form-check', output=output)
        receipt, _ = self.delivery._output(output['output_id'])
        result['items'] = [{'slide_id': 's1', 'verdict': 'passed', 'findings': [],
            'inspected_preview_sha256': receipt['files']['previews/s1.png'],
            'approved_preview_sha256': self.store.read()['slides']['s1']['build']['artifacts']['preview']['sha256']}]
        return task, result

    def test_export_requires_real_workflow_lock_and_current_revision(self):
        with self.assertRaisesRegex(ps.StateError, 'locked previews'):
            self.delivery.export(self.rev())
        self.assertEqual(self.assembler.call_count, 0)
        old = self.rev(); self.apply('set-content', payload=content())
        with self.assertRaisesRegex(ps.StateError, 'Revision conflict'):
            self.delivery.export(old)

    def test_export_default_does_not_prepare_optional_tasks_and_survives_temp_removal(self):
        with patch.object(self.delivery, 'prepare', side_effect=AssertionError('must not run')):
            output = self.exported()
        self.assertEqual(output['assistance'], pd._unchecked())
        self.assertEqual(output['source_policy'], 'local_only')
        self.assertFalse((self.store.root / 'requests').exists())
        shutil.rmtree(self.repo / 'temp')
        self.assertTrue(Path(output['artifacts']['pptx']).is_file())
        self.assertEqual(self.delivery.status()['status'], 'ready')
        receipt, folder = self.delivery._output(output['output_id'])
        self.assertEqual(len(ps.load_json(folder / 'deck-content.json')), 3)
        self.assertEqual(receipt['user_approval'], 'not_requested')

    def test_trigger_and_scope_are_explicit(self):
        output = self.exported()
        for kind, instruction, slides, question in [
                ('form-check', '', ['s1'], None), ('evidence-check', DECLARATION, ['s1'], None),
                ('form-check', DECLARATION, [], None), ('form-check', DECLARATION, ['s1', 's1'], None),
                ('form-check', DECLARATION, ['s9'], None), ('citation-redact', DECLARATION, ['s1'], 'extra')]:
            with self.subTest(kind=kind, slides=slides, question=question):
                with self.assertRaises(ps.StateError):
                    self.delivery.prepare(kind, self.rev(), slides, instruction, output_id=output['output_id'], question=question)

    def test_form_result_is_bound_scoped_and_not_a_user_approval(self):
        output = self.exported(); _, result = self.form(output)
        pointer = (self.store.root / 'current.json').read_bytes()
        for field, value in [('slide_id', 's2'), ('inspected_preview_sha256', '0' * 64), ('verdict', 'certified')]:
            bad = deepcopy(result); bad['items'][0][field] = value
            with self.assertRaises(ps.StateError):
                self.delivery.apply(bad)
        before = self.assembler.call_count
        self.assertEqual(self.delivery.apply(result, dry_run=True)['status'], 'validated')
        report = self.delivery.apply(result)
        self.assertEqual(self.assembler.call_count, before)
        self.assertEqual(report['assistance'], {**pd._unchecked(), 'form-check': 'executed'})
        self.assertEqual(report['user_approval'], 'not_requested')
        self.assertEqual(pointer, (self.store.root / 'current.json').read_bytes())
        with self.assertRaisesRegex(ps.StateError, 'already committed'):
            self.delivery.apply(result)

    def test_form_failure_and_unchecked_are_not_passed(self):
        output = self.exported(); _, result = self.form(output)
        result['items'][0].update(verdict='issues_found', findings=[])
        with self.assertRaises(ps.StateError): self.delivery.apply(result)
        result['items'][0]['findings'] = ['TEST-ONLY overflow finding']
        report = self.delivery.apply(result)
        stored = ps.load_json(Path(report['artifacts']['report']))
        self.assertEqual(stored['items'][0]['verdict'], 'issues_found')
        self.assertEqual(output['assistance'], pd._unchecked())

    def test_evidence_without_export_exact_quote_and_no_semantic_certification(self):
        path, c = self.source(); path.write_text('First line\nSynthetic result is 42.\n')
        c['sources'][0]['sha256'] = render.file_hash(path)
        self.apply('set-content', payload=c)
        task, result = self.candidate('evidence-check', question='TEST-ONLY: where does 42 come from?')
        self.assertEqual(task['schema'], 'agent-task-v1')
        result['items'] = [{'slide_id': 's1', 'claim': 'Synthetic result is 42', 'verdict': 'supports',
            'reason': 'TEST-ONLY semantic assertion', 'evidence': [
                {'source_id': 'src', 'locator': 'lines:2-2', 'quote': 'Synthetic result is 42.'}]}]
        for key, value in [('locator', 'lines:1-1'), ('quote', 'fabricated quote'), ('source_id', 'unknown')]:
            bad = deepcopy(result); bad['items'][0]['evidence'][0][key] = value
            with self.assertRaises(ps.StateError): self.delivery.apply(bad)
        check = self.delivery.apply(result, dry_run=True)
        self.assertEqual(check['semantic_verification'], 'not_performed_by_validator')
        self.delivery.apply(result)
        self.assertEqual(self.assembler.call_count, 0)
        self.assertEqual(self.store.read()['slides']['s1']['content']['sources'][0]['evidence_status'], 'user_provided')

    def test_evidence_insufficient_and_source_drift(self):
        path, c = self.source(); self.apply('set-content', payload=c)
        _, result = self.candidate('evidence-check', question='TEST-ONLY unsupported claim?')
        result['items'] = [{'slide_id': 's1', 'claim': 'test', 'verdict': 'supports', 'reason': 'test', 'evidence': []}]
        with self.assertRaises(ps.StateError): self.delivery.apply(result)
        result['items'][0]['verdict'] = 'insufficient'
        self.delivery.apply(result, dry_run=True)
        path.write_text('changed source')
        with self.assertRaisesRegex(ps.StateError, 'Source version changed'):
            self.delivery.apply(result)

    def test_redaction_is_separate_derivative_only_and_check_does_not_render(self):
        output = self.exported(); _, result = self.candidate('citation-redact', output=output)
        result['items'] = [{'slide_id': 's1', 'field': 'notes', 'before': 'Synthetic fixture only',
                            'after': 'Public-looking citation; still TEST-ONLY/local_only'}]
        before = self.assembler.call_count
        original = Path(output['artifacts']['pptx']).read_bytes()
        state = self.store.read()
        self.delivery.apply(result, dry_run=True)
        self.assertEqual(before, self.assembler.call_count)
        for field, value in [('slide_id', 's2'), ('before', 'wrong old text'), ('field', 'sources')]:
            bad = deepcopy(result); bad['items'][0][field] = value
            with self.assertRaises(ps.StateError): self.delivery.apply(bad)
        derived = self.delivery.apply(result)
        self.assertEqual(self.store.read(), state)
        self.assertEqual(Path(output['artifacts']['pptx']).read_bytes(), original)
        self.assertEqual(derived['assistance'], {**pd._unchecked(), 'citation-redact': 'executed'})
        receipt, folder = self.delivery._output(derived['output_id'])
        self.assertEqual(ps.load_json(folder / 'deck-content.json')[0]['content']['notes'], result['items'][0]['after'])
        self.assertTrue(receipt['limitations'])
        self.assertEqual(receipt['source_policy'], 'local_only')

    def test_corrupt_artifacts_request_and_symlink_are_rejected(self):
        output = self.exported(); _, result = self.form(output)
        request_dir = self.store.root / 'requests' / result['request_id']
        original = (request_dir / 'intent.json').read_bytes()
        (request_dir / 'intent.json').write_text('{}')
        with self.assertRaisesRegex(ps.StateError, 'integrity'):
            self.delivery.apply(result)
        (request_dir / 'intent.json').write_bytes(original)
        pptx = Path(output['artifacts']['pptx']); pptx.write_text('external edit')
        with self.assertRaisesRegex(ps.StateError, 'integrity'):
            self.delivery.apply(result)
        self.assertEqual(self.delivery.status()['status'], 'blocked')
        pptx.unlink(); pptx.symlink_to(self.repo / 'projects/demo')
        with self.assertRaisesRegex(ps.StateError, 'symlink'):
            self.delivery._output(output['output_id'])

    def test_source_and_renderer_changes_block_export(self):
        self.exported()
        with patch.object(render, 'fingerprint', return_value={'changed': True}):
            with self.assertRaisesRegex(ps.StateError, 'Renderer/font'):
                self.delivery.export(self.rev())
        root = self.store.root / 'outputs'; before = {p.name for p in root.iterdir()}
        with patch.object(render, 'assemble', side_effect=RuntimeError('renderer failed')):
            with self.assertRaises(RuntimeError): self.delivery.export(self.rev())
        self.assertEqual(before, {p.name for p in root.iterdir()})

    def test_stale_request_cannot_commit_after_revision_change(self):
        self.apply('set-content', payload=content())
        _, result = self.candidate('evidence-check', question='TEST-ONLY question')
        result['items'] = [{'slide_id': 's1', 'claim': 'test', 'verdict': 'insufficient', 'reason': 'test', 'evidence': []}]
        self.apply('set-content', payload=content(title='new'))
        with self.assertRaisesRegex(ps.StateError, 'Revision conflict'):
            self.delivery.apply(result)

    def test_two_processes_same_request_only_one_commit(self):
        output = self.exported(); _, result = self.form(output)
        ctx = mp.get_context('fork'); queue = ctx.Queue(); gate = ctx.Event()
        children = [ctx.Process(target=child_commit, args=(self.repo, self.store.root, result, queue, gate)) for _ in range(2)]
        for child in children: child.start()
        gate.set()
        messages = [queue.get(timeout=30) for _ in children]
        for child in children: child.join(30); self.assertEqual(child.exitcode, 0)
        self.assertEqual(messages.count('ok'), 1)
        self.assertTrue(any('already committed' in x for x in messages))

    def test_process_death_never_exposes_partial_output(self):
        output = self.exported()
        for checkpoint in ('file:result.json', 'before_publish', 'after_publish'):
            _, result = self.form(output)
            child = mp.get_context('fork').Process(target=child_commit,
                       args=(self.repo, self.store.root, result, None, None, checkpoint))
            child.start(); child.join(30); self.assertEqual(child.exitcode, 73)
            status = self.delivery.status(); self.assertEqual(status['status'], 'ready')
            committed = any(x['output_id'] == 'action-' + result['request_id'] for x in status['outputs'])
            self.assertEqual(committed, checkpoint == 'after_publish')
            if not committed: self.delivery.apply(result)
        self.assertTrue(self.delivery.status()['pending_directories'])

    def test_output_root_symlink_cannot_escape(self):
        outside = self.repo / 'outside'; outside.mkdir()
        (self.store.root / 'outputs').symlink_to(outside, target_is_directory=True)
        with self.assertRaises(ps.StateError): self.exported()
        self.assertEqual(list(outside.iterdir()), [])

    def test_result_schema_request_hash_and_duplicate_fields(self):
        output = self.exported(); _, result = self.form(output)
        for field, value in [('schema', 'future'), ('intent_sha256', '0' * 64), ('actor', '')]:
            bad = deepcopy(result); bad[field] = value
            with self.assertRaises(ps.StateError): self.delivery.apply(bad)
        bad = deepcopy(result); bad['extra'] = True
        with self.assertRaises(ps.StateError): self.delivery.apply(bad)
        bad = deepcopy(result); bad['items'] *= 2
        with self.assertRaises(ps.StateError): self.delivery.apply(bad)
        _, redaction = self.candidate('citation-redact', output=output)
        redaction['items'] = [{'slide_id': 's1', 'field': 'notes', 'before': 'Synthetic fixture only', 'after': 'test'}] * 2
        with self.assertRaises(ps.StateError): self.delivery.apply(redaction)
        redaction['items'] = redaction['items'][:1]
        with patch.object(render, 'fingerprint', return_value={'changed': True}):
            with self.assertRaisesRegex(ps.StateError, 'renderer changed'):
                self.delivery.apply(redaction, dry_run=True)

    def test_task_for_derivative_uses_actual_derivative_text(self):
        output = self.exported(); _, result = self.candidate('citation-redact', output=output)
        result['items'] = [{'slide_id': 's1', 'field': 'notes', 'before': 'Synthetic fixture only', 'after': 'New citation TEST-ONLY'}]
        derivative = self.delivery.apply(result)
        task, _ = self.candidate('citation-redact', output=derivative)
        context = next(x for x in task['inputs'] if x['name'] == 'context')
        self.assertEqual(ps.load_json(self.repo / context['path'])['slides']['s1']['notes'], 'New citation TEST-ONLY')

    def test_explicit_routes_and_real_cli_status(self):
        for kind in pd.KINDS:
            proc = subprocess.run([sys.executable, str(REPO / '.scripts/route.py'), '--capability', 'presentation',
                                   '--capability-profile', kind], capture_output=True, text=True)
            self.assertEqual(proc.returncode, 0, proc.stderr)
            self.assertIn(pd.LABELS[kind], proc.stdout)
            self.assertNotIn('## 阶段1A 命令与数据协议', proc.stdout)
        scripts = self.repo / '.scripts'; scripts.mkdir()
        for name in ('presentation_state.py', 'presentation_render.py', 'presentation_delivery.py', 'agent_task.py'):
            shutil.copy2(REPO / '.scripts' / name, scripts / name)
        proc = subprocess.run([sys.executable, str(scripts / 'presentation_delivery.py'), 'status', '--store', str(self.store.root)],
                              capture_output=True, text=True)
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        self.assertEqual(__import__('json').loads(proc.stdout)['outputs'], [])
        self.apply('set-content', payload=content())
        command = [sys.executable, str(scripts / 'presentation_delivery.py'), 'prepare', '--store', str(self.store.root),
                   '--expected-revision', self.rev(), '--kind', 'evidence-check', '--slides', 's1',
                   '--instruction', DECLARATION, '--question', 'TEST-ONLY: what is the source?']
        proc = subprocess.run(command, capture_output=True, text=True, cwd=self.repo)
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        task = __import__('json').loads(proc.stdout)
        candidate = self.repo / task['outputs'][0]['path']
        result = ps.load_json(candidate)
        result.update(actor='TEST-ONLY CLI', summary=DECLARATION, items=[{'slide_id': 's1', 'claim': 'Knowledge',
                      'verdict': 'insufficient', 'reason': 'TEST-ONLY no sources', 'evidence': []}])
        candidate.write_bytes(ps.json_bytes(result))
        import shlex
        for action in ('check', 'commit'):
            proc = subprocess.run(shlex.split(task['commands'][action]), capture_output=True, text=True, cwd=self.repo)
            self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        self.assertEqual(self.delivery.status()['outputs'][0]['kind'], 'evidence-check')


class RealDeliveryTests(unittest.TestCase):
    def test_real_three_page_export_and_text_citation_derivative(self):
        from PIL import Image, ImageDraw
        with tempfile.TemporaryDirectory(prefix='presentation-delivery-real-') as temporary:
            repo = Path(temporary).resolve(); project = repo / 'projects/fixture'; project.mkdir(parents=True)
            image = Image.new('RGB', (600, 400), '#E0EBFA')
            ImageDraw.Draw(image).rectangle((100, 90, 480, 300), fill='#3265A8')
            material = project / 'synthetic.png'; image.save(material)
            import fitz
            pdf_source = project / 'synthetic-reference.pdf'
            with fitz.open() as document:
                document.new_page().insert_text((72, 72), 'Synthetic reference value is 42.')
                document.save(pdf_source)
            store = ps.Store(project / 'presentations/fixture', repo); store.initialize(ps.example_plan())
            for sid, layout in [('s1', 'title'), ('s2', 'image_text'), ('s3', 'flow')]:
                c = content(sid, '合成测试 TEST-ONLY · ' + layout)
                c['body'] = ['Knowledge 知识', 'Evidence 证据']
                c['notes'] = 'Source: INTERNAL-CITATION-TEST-ONLY'
                c['sources'] = [{'id': 'reference', 'path': pdf_source.relative_to(repo).as_posix(),
                    'sha256': render.file_hash(pdf_source), 'locator': 'page:1',
                    'evidence_status': 'user_provided', 'sensitivity': 'public'}]
                if layout == 'image_text':
                    c['sources'] += [{'id': 'image', 'path': material.relative_to(repo).as_posix(),
                        'sha256': render.file_hash(material), 'locator': 'synthetic fixture',
                        'evidence_status': 'user_provided', 'sensitivity': 'public'}]
                for op, value in [('set-content', c)]: store.apply(request(store, op, sid, value))
                store.apply(request(store, 'approve-content', sid, approval(store, sid, 'content')))
                store.apply(request(store, 'set-design', sid, design(sid, layout, 'image' if layout == 'image_text' else None)))
                store.apply(request(store, 'approve-design', sid, approval(store, sid, 'design')))
                store.apply(request(store, 'build', sid, {}))
                # Explicitly TEST-ONLY state fixtures; not a real user or Agent visual acceptance.
                store.apply(request(store, 'review', sid, approval(store, sid, 'visual')))
                store.apply(request(store, 'lock', sid, approval(store, sid, 'preview', store.read()['slides'][sid]['build']['warnings'])))
            delivery = pd.Delivery(store); rev = store.read()['session']['revision_id']
            output = delivery.export(rev)
            self.assertEqual(output['assistance'], pd._unchecked())
            from pptx import Presentation
            deck = Presentation(output['artifacts']['pptx'])
            self.assertEqual(len(deck.slides), 3)
            self.assertTrue(deck.slides[0].shapes[0].has_text_frame)
            self.assertEqual([page.shapes[0].text for page in deck.slides],
                             ['合成测试 TEST-ONLY · ' + layout for layout in ['title', 'image_text', 'flow']])
            deck.slides[0].shapes[0].text = 'Native text remains editable TEST-ONLY'
            deck.save(project / 'editing-test.pptx')
            self.assertEqual(Presentation(project / 'editing-test.pptx').slides[0].shapes[0].text, 'Native text remains editable TEST-ONLY')
            import fitz
            with fitz.open(output['artifacts']['pdf']) as pdf:
                self.assertEqual(len(pdf), 3)
                for page in pdf:
                    self.assertIn('Knowledge', page.get_text())
            evidence_task = delivery.prepare('evidence-check', rev, ['s1'], DECLARATION,
                                             question='TEST-ONLY: where does the reference value come from?')
            evidence_result = ps.load_json(repo / evidence_task['outputs'][0]['path'])
            evidence_result.update(actor='TEST-ONLY', summary=DECLARATION, items=[{'slide_id': 's1',
                'claim': 'Synthetic reference value is 42.', 'verdict': 'supports', 'reason': DECLARATION,
                'evidence': [{'source_id': 'reference', 'locator': 'page:1', 'quote': 'Synthetic reference value is 42.'}]}])
            bad = deepcopy(evidence_result); bad['items'][0]['evidence'][0]['locator'] = 'page:2'
            with self.assertRaises(ps.StateError): delivery.apply(bad)
            delivery.apply(evidence_result)
            task = delivery.prepare('citation-redact', rev, ['s1', 's2', 's3'], DECLARATION, output_id=output['output_id'])
            result = ps.load_json(repo / task['outputs'][0]['path'])
            result.update(actor='TEST-ONLY', summary=DECLARATION, items=[
                {'slide_id': sid, 'field': 'notes', 'before': 'Source: INTERNAL-CITATION-TEST-ONLY',
                 'after': 'Source: synthetic test fixture'} for sid in ['s1', 's2', 's3']])
            derivative = delivery.apply(result)
            import zipfile
            with zipfile.ZipFile(derivative['artifacts']['pptx']) as archive:
                self.assertFalse(any(b'INTERNAL-CITATION-TEST-ONLY' in archive.read(name) for name in archive.namelist() if name.endswith('.xml')))
            self.assertEqual(store.read()['session']['revision_id'], rev)
            saved_root = REPO / 'temp/presentation-runtime'; saved_root.mkdir(parents=True, exist_ok=True)
            saved = Path(tempfile.mkdtemp(prefix='delivery-three-pages-', dir=saved_root))
            for name, path in output['artifacts'].items(): shutil.copy2(path, saved / Path(path).name)
            shutil.copy2(derivative['artifacts']['pptx'], saved / 'citation-redacted.pptx')
            (saved / 'README.txt').write_text('Synthetic TEST-ONLY three-page mechanical export. Simulated fixture approvals are NOT real user acceptance. No automatic form/evidence review. PowerPoint GUI check remains pending.\n')
            print('REAL DELIVERY FIXTURE:', saved)


def main():
    parser = argparse.ArgumentParser(); parser.add_argument('--render', action='store_true'); args = parser.parse_args()
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(DeliveryTests)
    if args.render: suite.addTests(unittest.defaultTestLoader.loadTestsFromTestCase(RealDeliveryTests))
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    if not args.render: print('Real LibreOffice render not executed; use --render to require it.')
    return 0 if result.wasSuccessful() else 1


if __name__ == '__main__':
    raise SystemExit(main())
