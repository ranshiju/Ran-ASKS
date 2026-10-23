#!/usr/bin/env python3
"""Local native A1/A0 exhibit compiler; no semantic judgments or remote calls."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import re
import subprocess
import unicodedata

CANVAS_MM = [594, 841]
SUPPORTED_CANVASES = ((594, 841), (841, 1189))
POINTS_PER_MM = 72 / 25.4
LINE_START_PROHIBITED = '，。、；：！？）》】〕〉」』”’％%,.!?:;)]}'
LINE_END_PROHIBITED = '（《【〔〈「『“‘([{'


def typographic_units(tokens):
    units = []
    for token in tokens:
        if units and (token[0] in LINE_START_PROHIBITED or units[-1][-1] in LINE_END_PROHIBITED):
            units[-1] += token
        else:
            units.append(token)
    return units


def file_hash(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def compiler_hash():
    import presentation_artwork
    import science_exhibit
    import science_exhibit_composition
    import review_contract
    return hashlib.sha256((file_hash(__file__) + file_hash(presentation_artwork.__file__) + file_hash(science_exhibit.__file__) + file_hash(science_exhibit_composition.__file__) + file_hash(review_contract.__file__)).encode()).hexdigest()


def font_info(name):
    import visual_qa
    environment = visual_qa._soffice_env(visual_qa._find_soffice())
    result = subprocess.run(
        ['fc-match', '--format=%{family}\n%{file}\n%{index}\n', name],
        capture_output=True, text=True, check=True, timeout=30, env=environment,
    )
    lines = result.stdout.splitlines()
    if len(lines) < 3 or name.casefold() not in {family.strip().casefold() for family in lines[0].split(',')}:
        raise ValueError(f'Requested font unavailable: {name}')
    path = Path(lines[1])
    if not path.is_file():
        raise ValueError('Requested font file unavailable')
    return {'family': name, 'path': str(path), 'index': int(lines[2]), 'sha256': file_hash(path)}


def wrap_text(text, width_mm, height_mm, size_pt, font):
    from PIL import ImageFont
    scale = 4
    measured = ImageFont.truetype(font['path'], round(size_pt * scale), index=font['index'])
    width = width_mm * POINTS_PER_MM * scale * 0.97
    lines = []
    for paragraph in text.split('\n'):
        current = ''
        tokens = re.findall(r'[A-Za-z0-9][A-Za-z0-9_./+\-]*|[^\s]|[ \t]+', paragraph)
        for token in typographic_units(tokens):
            pieces = typographic_units(list(token)) if measured.getlength(token) > width else [token]
            for piece in pieces:
                if measured.getlength(piece) > width:
                    raise ValueError('Text box too narrow for a typographic unit')
                candidate = current + piece
                if current and measured.getlength(candidate) > width:
                    lines.append(current.rstrip())
                    current = piece.lstrip()
                else:
                    current = candidate
        lines.append(current.rstrip())
    if len(lines) * size_pt * 1.3 > height_mm * POINTS_PER_MM:
        raise ValueError('Text capacity exceeded; shorten the content or change the layout, not the font size')
    return '\n'.join(lines)


def compile_layout(content, brief, font, canvas_mm=None):
    import science_exhibit
    science_exhibit.validate_display_acronyms(science_exhibit.display_texts(content, brief))
    canvas = list(CANVAS_MM if canvas_mm is None else canvas_mm)
    if tuple(canvas) not in SUPPORTED_CANVASES:
        raise ValueError('Supported canvases are A1/A0 portrait only')
    if content.get('schema') == 'science-exhibit-content-v2':
        import science_exhibit_composition
        return science_exhibit_composition.compile_layout(content, brief, font, canvas, wrap_text)
    objects = []

    def add_text(identifier, text, box, size, color='18323B', bold=False, claim_ids=None, panel_id=None):
        objects.append({
            'id': identifier, 'type': 'text', 'bbox_mm': box, 'font_pt': size,
            'color': color, 'bold': bold,
            'claim_ids': list(claim_ids or []), 'panel_id': panel_id,
            'text': wrap_text(text, box[2], box[3], size, font),
        })

    add_text('draft-label', '科学展板 · 待审草稿', [28, 17, 538, 12], 22, '526B74')
    concept = content['diagram'].get('kind') == 'concept'
    add_text('title', content['title'], [28, 39, 538, 64], 64 if concept else 76, bold=True)
    add_text('hook', content['hook'], [28, 115, 538, 29], 34, '236F78')
    diagram = content['diagram']
    if concept:
        objects.append({'id': 'takeaway-band', 'type': 'shape', 'shape_kind': 'rect',
                        'bbox_mm': [28, 153, 538, 49], 'color': 'E4F0EE', 'stroke': 'none',
                        'stroke_pt': 0, 'radius_mm': 4, 'claim_ids': list(content['takeaway']['claim_ids'])})
    add_text('takeaway', content['takeaway']['text'], [28, 157, 538, 42 if concept else 60], 40, bold=True, claim_ids=content['takeaway']['claim_ids'])
    add_text('diagram-heading', diagram['headline'] + '（科学示意，非观测图）', [28, 224 if concept else 236, 538, 18], 25, '526B74')
    if concept:
        artwork = diagram['artwork']
        width, height = artwork['canvas']
        scale = min(538 / width, 310 / height)
        origin = [28 + (538 - width * scale) / 2, 248 + (310 - height * scale) / 2]
        objects.append({'id': 'artwork-background', 'type': 'shape', 'shape_kind': 'rect',
                        'bbox_mm': [*origin, width * scale, height * scale], 'color': artwork['background'][1:],
                        'stroke': 'none', 'stroke_pt': 0, 'radius_mm': 0, 'claim_ids': list(diagram['claim_ids'])})
        for index, primitive in enumerate(artwork['objects']):
            identifier = f'artwork-{index}'
            kind = primitive['type']
            if kind == 'line':
                objects.append({'id': identifier, 'type': 'line',
                                'start_mm': [origin[0] + primitive['x1'] * scale, origin[1] + primitive['y1'] * scale],
                                'end_mm': [origin[0] + primitive['x2'] * scale, origin[1] + primitive['y2'] * scale],
                                'stroke': primitive['stroke'][1:], 'stroke_pt': primitive['stroke_width'] * scale * POINTS_PER_MM,
                                'claim_ids': list(diagram['claim_ids'])})
                continue
            box = [origin[0] + primitive['x'] * scale, origin[1] + primitive['y'] * scale,
                   primitive['w'] * scale, primitive['h'] * scale]
            if kind == 'text':
                size = primitive['size'] * scale * POINTS_PER_MM
                if size < 24:
                    raise ValueError('Concept labels must be at least 24 pt; simplify the visual or enlarge the labels')
                add_text(identifier, primitive['text'], box, size, primitive['color'][1:], primitive['bold'], diagram['claim_ids'])
                objects[-1]['align'] = primitive['align']
            else:
                objects.append({'id': identifier, 'type': 'shape', 'shape_kind': kind, 'bbox_mm': box,
                                'color': primitive['fill'][1:] if primitive['fill'] != 'none' else 'none',
                                'stroke': primitive['stroke'][1:] if primitive['stroke'] != 'none' else 'none',
                                'stroke_pt': primitive['stroke_width'] * scale * POINTS_PER_MM,
                                'radius_mm': primitive.get('radius', 0) * scale, 'claim_ids': list(diagram['claim_ids'])})
    else:
        steps = diagram['steps']
        width = (538 - 12 * (len(steps) - 1)) / len(steps)
        for index, step in enumerate(steps):
            left = 28 + index * (width + 12)
            objects.append({'id': step['id'], 'type': 'shape', 'bbox_mm': [left, 268, width, 76], 'color': 'E4F0EE', 'claim_ids': list(step['claim_ids'])})
            add_text(f'{step["id"]}-text', step['text'], [left + 8, 280, width - 16, 53], 28, bold=True, claim_ids=step['claim_ids'])
            if index:
                objects.append({'id': f'connector-{index}', 'type': 'connector', 'start': steps[index - 1]['id'], 'end': step['id']})
    for index, panel in enumerate(content['panels']):
        left = 28 + (index % 2) * 281
        top = (584 + (index // 2) * 96) if concept else (380 + (index // 2) * 198)
        add_text(f'{panel["id"]}-headline', panel['headline'], [left, top, 257, 25], 34 if concept else 32, '236F78', True, panel['claim_ids'], panel['id'])
        add_text(f'{panel["id"]}-body', panel['body'], [left, top + 34 if concept else top + 36, 257, 62 if concept else 145], 30 if concept else 27, claim_ids=panel['claim_ids'], panel_id=panel['id'])
    add_text('source-label', '研究来源：' + brief['source_label'], [28, 786, 538, 26], 18, '526B74')
    add_text('review-label', '来源定位已检查；科学表达、视觉效果与公开展示授权仍待确认。', [28, 824, 538, 10], 16, '526B74')
    factor = min(actual / nominal for actual, nominal in zip(canvas, CANVAS_MM))
    offset = [(actual - nominal * factor) / 2 for actual, nominal in zip(canvas, CANVAS_MM)]
    for obj in objects:
        if 'bbox_mm' in obj:
            left, top, width, height = obj['bbox_mm']
            obj['bbox_mm'] = [left * factor + offset[0], top * factor + offset[1], width * factor, height * factor]
        for key in ('start_mm', 'end_mm'):
            if key in obj:
                obj[key] = [coordinate * factor + origin for coordinate, origin in zip(obj[key], offset)]
        for key in ('font_pt', 'stroke_pt', 'radius_mm'):
            if key in obj:
                obj[key] *= factor
    return {'schema': 'science-exhibit-layout-v1', 'canvas_mm': canvas, 'visual_kind': 'concept' if concept else 'flow', 'objects': objects}


def normalized(text):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', text))


def validate_rendered_lines(objects, rendered_lines):
    import fitz
    assigned = set()
    for obj in objects:
        if obj['type'] != 'text':
            continue
        left, top, width, height = obj['bbox_mm']
        rectangle = fitz.Rect(left, top, left + width, top + height) * POINTS_PER_MM
        rectangle = rectangle + (-0.7, -0.7, 0.7, 0.7)
        previous_top = -float('inf')
        for value in obj['text'].split('\n'):
            if not value.strip():
                continue
            candidates = [(index, box) for index, (box, actual) in enumerate(rendered_lines)
                          if index not in assigned and normalized(actual) == normalized(value) and
                          box.y0 >= previous_top - 0.7 and rectangle.contains(box)]
            if not candidates:
                raise ValueError(f'Unexpected rendered wrapping or clipped text: {obj["id"]}')
            index, box = min(candidates, key=lambda item: (item[1].y0, item[1].x0))
            assigned.add(index)
            previous_top = box.y0
    if any(index not in assigned and normalized(value) for index, (_, value) in enumerate(rendered_lines)):
        raise ValueError('PDF contains unexpected or duplicated text lines')


def render(content, brief, directory, *, canvas_mm=None):
    compiler_sha256 = compiler_hash()
    import presentation_runtime
    presentation_runtime._load_pptx()
    import fitz
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Mm, Pt
    import visual_qa
    import science_exhibit
    import science_exhibit_composition
    import io

    directory = Path(directory)
    font = font_info(brief['font'])
    layout = compile_layout(content, brief, font, canvas_mm)
    assets = science_exhibit_composition.asset_paths(science_exhibit, content) if content.get('schema') == 'science-exhibit-content-v2' else {}
    asset_data = {identifier: path.read_bytes() for identifier, path in assets.items()}
    expected_assets = {item['id']: item for item in content.get('assets', [])}
    for identifier, data in asset_data.items():
        if hashlib.sha256(data).hexdigest() != expected_assets[identifier]['sha256']:
            raise ValueError('Asset changed during rendering')
    canvas = layout['canvas_mm']
    deck = Presentation()
    deck.slide_width, deck.slide_height = [Mm(value) for value in canvas]
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(layout.get('background', 'F7FAF8'))
    native = {}
    for obj in layout['objects']:
        if obj['type'] == 'connector':
            continue
        if obj['type'] == 'image':
            shape = slide.shapes.add_picture(io.BytesIO(asset_data[obj['asset_id']]), *[Mm(value) for value in obj['bbox_mm']])
            shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom = obj.get('crop', [0, 0, 0, 0])
        elif obj['type'] == 'line':
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *[Mm(value) for value in [*obj['start_mm'], *obj['end_mm']]])
            shape.line.color.rgb = RGBColor.from_string(obj['stroke'])
            shape.line.width = Pt(obj['stroke_pt'])
        elif obj['type'] == 'shape':
            kind = obj.get('shape_kind', 'rect')
            radius = obj.get('radius_mm', 1)
            shape_type = MSO_SHAPE.OVAL if kind == 'ellipse' else MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
            shape = slide.shapes.add_shape(shape_type, *[Mm(value) for value in obj['bbox_mm']])
            if kind == 'rect' and 'radius_mm' in obj and radius:
                shape.adjustments[0] = radius / min(obj['bbox_mm'][2:])
            if obj['color'] == 'none':
                shape.fill.background()
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(obj['color'])
            if obj.get('stroke', 'none') == 'none' or obj.get('stroke_pt', 0) == 0:
                shape.line.fill.background()
            else:
                shape.line.color.rgb = RGBColor.from_string(obj['stroke'])
                shape.line.width = Pt(obj['stroke_pt'])
        else:
            shape = slide.shapes.add_textbox(*[Mm(value) for value in obj['bbox_mm']])
            frame = shape.text_frame
            frame.word_wrap = False
            frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
            for index, line in enumerate(obj['text'].split('\n')):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.line_spacing = Pt(obj['font_pt'] * 1.3)
                paragraph.space_before = paragraph.space_after = Pt(0)
                paragraph.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[obj.get('align', 'left')]
                run = paragraph.add_run()
                run.text = line
                run.font.name = font['family']
                run.font.size = Pt(obj['font_pt'])
                run.font.bold = obj['bold']
                run.font.color.rgb = RGBColor.from_string(obj['color'])
                east_asian = OxmlElement('a:ea')
                east_asian.set('typeface', font['family'])
                run._r.get_or_add_rPr().append(east_asian)
        style = shape.element.find(qn('p:style'))
        if style is not None:
            shape.element.remove(style)
        shape.name = obj['id']
        native[obj['id']] = shape
    for obj in layout['objects']:
        if obj['type'] != 'connector':
            continue
        start, end = native[obj['start']], native[obj['end']]
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, start.left + start.width,
            start.top + start.height // 2, end.left, end.top + end.height // 2,
        )
        connector.begin_connect(start, 3)
        connector.end_connect(end, 1)
        connector.line.color.rgb = RGBColor.from_string('236F78')
        connector.line.width = Pt(2 * min(actual / nominal for actual, nominal in zip(canvas, CANVAS_MM)))
        style = connector.element.find(qn('p:style'))
        if style is not None:
            connector.element.remove(style)
        connector.name = obj['id']
    pptx = directory / 'board.pptx'
    deck.save(pptx)
    reopened = Presentation(pptx)
    if len(reopened.slides) != 1 or len(reopened.slides[0].shapes) != len(layout['objects']):
        raise ValueError('Native object readback failed')
    if any(abs(actual.mm - expected) > 0.001 for actual, expected in zip((reopened.slide_width, reopened.slide_height), canvas)):
        raise ValueError('PPTX physical size mismatch')
    actual_text = {shape.name: shape.text for shape in reopened.slides[0].shapes if shape.has_text_frame}
    by_name = {shape.name: shape for shape in reopened.slides[0].shapes}
    for obj in layout['objects']:
        if obj['type'] == 'line':
            shape = by_name[obj['id']]
            actual_endpoints = [dimension.mm for dimension in (shape.begin_x, shape.begin_y, shape.end_x, shape.end_y)]
            if any(abs(actual - expected) > 0.001 for actual, expected in zip(actual_endpoints, [*obj['start_mm'], *obj['end_mm']])):
                raise ValueError(f'Native line readback failed: {obj["id"]}')
        if 'bbox_mm' in obj:
            shape = by_name[obj['id']]
            actual_box = [dimension.mm for dimension in (shape.left, shape.top, shape.width, shape.height)]
            if any(abs(actual - expected) > 0.001 for actual, expected in zip(actual_box, obj['bbox_mm'])):
                raise ValueError(f'Native geometry readback failed: {obj["id"]}')
        if obj['type'] == 'text' and actual_text.get(obj['id']) != obj['text']:
            raise ValueError(f'Native text readback failed: {obj["id"]}')
        if obj['type'] == 'image' and hashlib.sha256(by_name[obj['id']].image.blob).hexdigest() != expected_assets[obj['asset_id']]['sha256']:
            raise ValueError(f'Embedded image readback failed: {obj["id"]}')
        if obj['type'] == 'image' and any(abs(actual - expected) > 0.00002 for actual, expected in zip(
                (by_name[obj['id']].crop_left, by_name[obj['id']].crop_top, by_name[obj['id']].crop_right, by_name[obj['id']].crop_bottom),
                obj.get('crop', [0, 0, 0, 0]))):
            raise ValueError(f'Image crop readback failed: {obj["id"]}')
        if obj['type'] == 'connector':
            for tag, endpoint in (('stCxn', 'start'), ('endCxn', 'end')):
                binding = by_name[obj['id']].element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}' + tag)
                if binding is None or binding.get('id') != str(by_name[obj[endpoint]].shape_id):
                    raise ValueError('Native connector binding readback failed')
    pdf = directory / 'board.pdf'
    visual_qa.convert_slides_to_pdf(pptx, pdf)
    with fitz.open(pdf) as document:
        if document.page_count != 1:
            raise ValueError('PDF page count mismatch')
        page = document[0]
        if any(abs(actual - expected * POINTS_PER_MM) > 0.5 for actual, expected in zip((page.rect.width, page.rect.height), canvas)):
            raise ValueError('PDF physical size mismatch')
        text = normalized(page.get_text())
        for obj in layout['objects']:
            if obj['type'] == 'text' and normalized(obj['text']) not in text:
                raise ValueError(f'PDF lost text or contains missing glyphs: {obj["id"]}')
        fonts = page.get_fonts(full=True)
        requested = normalized(font['family']).casefold()
        if not fonts or any(requested not in normalized(record[3]).casefold() for record in fonts):
            raise ValueError('PDF font substitution detected')
        if any(not document.extract_font(record[0])[3] for record in fonts):
            raise ValueError('PDF font not embedded')
        rendered_lines = []
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines', []):
                if not page.rect.contains(fitz.Rect(line['bbox'])):
                    raise ValueError('Rendered text outside the page')
                rendered_lines.append((fitz.Rect(line['bbox']), ''.join(span['text'] for span in line['spans'])))
        validate_rendered_lines(layout['objects'], rendered_lines)
        page.get_pixmap(matrix=fitz.Matrix(0.6, 0.6), alpha=False).save(directory / 'preview.png')
        font_names = sorted({record[3] for record in fonts})
    if file_hash(font['path']) != font['sha256']:
        raise ValueError('Font changed during rendering')
    if compiler_hash() != compiler_sha256:
        raise ValueError('Compiler changed during rendering')
    if assets:
        science_exhibit_composition.asset_paths(science_exhibit, content)
    (directory / 'layout.json').write_text(json.dumps(layout, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return {
        'schema': 'science-exhibit-render-v1', 'mechanical_status': 'passed',
        'canvas_mm': canvas, 'native_object_count': len(layout['objects']),
        'requested_font': font['family'], 'font_sha256': font['sha256'], 'pdf_fonts': font_names,
        'renderer_sha256': compiler_sha256,
        'image_object_count': sum(obj['type'] == 'image' for obj in layout['objects']),
        'image_assets': {identifier: {'sha256': item['sha256'], 'width': item['width'], 'height': item['height']}
                         for identifier, item in expected_assets.items()},
        'image_effective_dpi': {obj['id']: round(min(expected_assets[obj['asset_id']]['width'] * (1 - obj.get('crop', [0, 0, 0, 0])[0] - obj.get('crop', [0, 0, 0, 0])[2]) / (obj['bbox_mm'][2] / 25.4),
                                                    expected_assets[obj['asset_id']]['height'] * (1 - obj.get('crop', [0, 0, 0, 0])[1] - obj.get('crop', [0, 0, 0, 0])[3]) / (obj['bbox_mm'][3] / 25.4)), 1)
                                for obj in layout['objects'] if obj['type'] == 'image'},
        'visual_review': 'not_checked', 'semantic_review': 'pending',
        'public_release': 'not_approved', 'remote_calls': 0,
    }
